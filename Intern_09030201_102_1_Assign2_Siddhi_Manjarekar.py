import random
import unicodedata
from matplotlib.ticker import MultipleLocator
import matplotlib.font_manager as fm
import openpyxl
import warnings
import base64
from io import BytesIO
import os

os.environ['MPLBACKEND'] = 'Agg'
os.environ['MPLCONFIGDIR'] = os.getcwd()

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

plt.rcParams['text.usetex'] = False
plt.rcParams['svg.fonttype'] = 'none'
plt.rcParams['text.hinting'] = 'none'
plt.rcParams['text.antialiased'] = True

warnings.filterwarnings('ignore')


def setup_marathi_font():
    font_files = [
        "NotoSansDevanagari-Regular.ttf",
        "NotoSansDevanagari-Bold.ttf",
        "Lohit-Devanagari.ttf",
        "Mangal.ttf",
        "Kokila.ttf"
    ]

    for font_file in font_files:
        if os.path.exists(font_file):
            try:
                fm.fontManager.addfont(font_file)
                font_prop = fm.FontProperties(fname=font_file)
                plt.rcParams['font.family'] = font_prop.get_name()
                return font_prop
            except Exception:
                continue

    system_fonts = fm.findSystemFonts(fontext='ttf')
    devanagari_fonts = [f for f in system_fonts if any(x in f.lower() 
                        for x in ['devanagari', 'noto', 'mangal', 'lohit', 'kokila', 'sahadeva'])]
    
    for fp in devanagari_fonts:
        try:
            fm.fontManager.addfont(fp)
            font_prop = fm.FontProperties(fname=fp)
            plt.rcParams['font.family'] = font_prop.get_name()
            return font_prop
        except Exception:
            continue

    return fm.FontProperties()


def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)


def create_chart_base64(rainfall, days, question_no=None):
    font_prop = setup_marathi_font()
    
    fig = plt.figure(figsize=(5.5, 3.5), dpi=90, facecolor='white')
    ax = fig.add_axes([0.15, 0.15, 0.80, 0.75])
    
    ax.set_facecolor('white')

    bar_colors = ['#4472C4', '#ED7D31', '#A5A5A5', '#FFC000', 
                  '#5B9BD5', '#70AD47', '#264478']
    
    bars = ax.bar(range(len(days)), rainfall, color=bar_colors, 
                  edgecolor='black', linewidth=0.7, zorder=3, width=0.75)

    y_max = max(50, int(max(rainfall) + 5))
    ax.set_ylim(0, y_max)
    ax.yaxis.set_major_locator(MultipleLocator(5))
    ax.yaxis.set_minor_locator(MultipleLocator(1))

    ax.grid(which='major', axis='y', linestyle='-', linewidth=1.0, 
            color='gray', alpha=0.8, zorder=0)
    ax.grid(which='minor', axis='y', linestyle='-', linewidth=0.5, 
            color='gray', alpha=0.6, zorder=0)
    
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_linewidth(1.0)
    ax.spines['bottom'].set_linewidth(1.0)

    ax.set_xticks(range(len(days)))
    ax.set_xticklabels([''] * len(days))  
    
    for label in ax.get_yticklabels():
        label.set_fontproperties(font_prop)
        label.set_fontsize(8)

    try:
        buffer = BytesIO()
        fig.savefig(buffer, format='png', dpi=95, bbox_inches='tight',
                    facecolor='white', edgecolor='none', pil_kwargs={'optimize': True})
        buffer.seek(0)
        img_data = buffer.read()
        b64_str = base64.b64encode(img_data).decode('ascii')
        plt.close(fig)
        return f"data:image/png;base64,{b64_str}"
    except Exception as e:
        print(f"Error generating chart: {e}")
        plt.close(fig)
        return ""


def generate_wrong_answers(rainfall, correct_ans_values):
    wrong_answers_set = set()
    attempts = 0
    max_attempts = 50
    
    while len(wrong_answers_set) < 3 and attempts < max_attempts:
        attempts += 1
        temp_rainfall = rainfall.copy()
        method = random.choice(['change', 'shuffle', 'swap'])

        if method == 'change':
            idx = random.randrange(len(temp_rainfall))
            change = random.choice([-5, 5, -4, 4, -3, 3, -2, 2])
            temp_rainfall[idx] = max(0, min(50, temp_rainfall[idx] + change))
        elif method == 'shuffle':
            random.shuffle(temp_rainfall)
            if tuple(temp_rainfall) == tuple(rainfall):
                continue
        elif method == 'swap':
            idx1, idx2 = random.sample(range(len(temp_rainfall)), 2)
            temp_rainfall[idx1], temp_rainfall[idx2] = temp_rainfall[idx2], temp_rainfall[idx1]

        wrong_ans_str = f"${',\\\\ '.join(map(str, temp_rainfall))}$"
        if wrong_ans_str != correct_ans_values:
            wrong_answers_set.add(wrong_ans_str)

    return list(wrong_answers_set) if len(wrong_answers_set) == 3 else None


def create_solution_text(rainfall, correct_option_letter, scale=5):
    eng_sol = (
        f"Ans : Option ${correct_option_letter}$.<br>"
        f"Bar graph shows daily rainfall in mm, from Monday to Sunday.<br>"
        f"Corresponding to each day, there is $1$ bar as shown.<br>"
        f"Scale is mentioned in the graph as $1$ unit $= {scale}$ mm of rainfall.<br>"
        f"Each unit is further divided in to $5$ smaller divisions.<br>"
        f"Hence, each smaller division $= \\dfrac {{1}}{{5}}= 0.2$ unit.<br>"
        f"Quantity of rainfall can be calculated by knowing the height of the bar for each day.<br>"
        f"Formula used for this is,<br>"
        f"Daily rainfall $ = $ height of the bar $\\times$ scale used.<br>"
        f"For Monday, height of the bar is {rainfall[0]/scale:.1f} units,<br>"
        f"Hence, rainfall on Monday $ = {rainfall[0]/scale:.1f} \\times {scale} = {rainfall[0]}$ mm.<br>"
        f"Similarly we can calculate the rainfall for Tuesday.<br>"
        f"For Tuesday, height of the bar is {rainfall[1]/scale:.1f} units,<br>"
        f"Hence, rainfall on Tuesday $ = {rainfall[1]/scale:.1f} \\times {scale} = {rainfall[1]}$ mm.<br>"
        f"Rainfall for remaining days can be calculated similarly to get the values of rainfall as,<br>"
        f"Wednesday ${rainfall[2]}$ mm, Thursday ${rainfall[3]}$ mm, Friday ${rainfall[4]}$ mm, "
        f"Saturday ${rainfall[5]}$ mm and Sunday ${rainfall[6]}$ mm of rainfall.<br>"
        f"Hence, the rainfall in the week from Monday to Sunday is ${',\\\\ '.join(map(str, rainfall))}$ mm<br>"
        f"By comparing these values with the values in different options sequentially, we come to know that, "
        f"these values match only with the values in option ${correct_option_letter}$ perfectly.<br>"
        f"In remaining options, at least one of the value does not match with these values.<br>"
        f"Hence, option ${correct_option_letter}$, is the answer.<br>#"
    )

    mar_sol = (
        f"उत्तर : पर्याय ${correct_option_letter}$.<br>"
        f"स्तंभालेख सप्ताहात सोमवार ते रविवार रोज पडणार्‍या, पावसाचे मान (मोजणी) म्हणजेच पर्जन्य मान मिमी या एककात दाखवितो.<br>"
        f"प्रत्येक दिवसासाठी एक या प्रमाणे या स्तंभालेखात $7$ स्तंभ आहेत.<br>"
        f"स्तंभालेखात $Y$ अक्षावरील प्रमाण $1$ एकक $= {scale}$ मिमी पाऊस असे आहे.<br>"
        f"प्रत्येक एककाचे घर पुन्हा $5$ लहान घरात विभागले आहे.<br>"
        f"म्हणून, प्रत्येक लहान घर $= \\dfrac {{1}}{{5}}= 0.2$ एकक.<br>"
        f"आलेखातील प्रत्येक स्तंभाच्या उंची नुसार त्या दिवसा साठीचा पडलेला पाऊस शोधात येईल.<br>"
        f"या साठी वापरायचे सूत्र असे,<br>"
        f"दैनंदिन पाऊस $ = $ त्या दिवसा साठीची स्तंभाची उंची $\\times$ प्रमाण.<br>"
        f"सोमवार साठी स्तंभाची उंची ${rainfall[0]/scale:.1f}$ एकक आहे. <br>"
        f"म्हणून सोमवारी पडलेला पाऊस $ = {rainfall[0]/scale:.1f} \\times {scale} = {rainfall[0]}$ मिमी.<br>"
        f"याच पद्धतीने मंगळवारी पडलेला पाऊस ठरवता येऊ शकतो.<br>"
        f"मंगळवार साठी स्तंभाची उंची ${rainfall[1]/scale:.1f}$ एकक आहे,<br>"
        f"म्हणून मंगळवारी पडलेला पाऊस $ = {rainfall[1]/scale:.1f} \\times {scale} = {rainfall[1]}$ मिमी.<br>"
        f"याच पद्धतीने बुधवार ${rainfall[2]}$ मिमी, गुरुवार ${rainfall[3]}$ मिमी, शुक्रवार ${rainfall[4]}$ मिमी, "
        f"शनिवार ${rainfall[5]}$ मिमी आणि रविवार ${rainfall[6]}$ मिमी पाऊस पडला असे समजते.<br>"
        f"म्हणून, सप्ताहातील पाऊस ${',\\\\ '.join(map(str, rainfall))}$ मिमी पाऊस पडला. <br>"
        f"दिलेल्या प्रत्येक पर्यायाशी या संख्यांची क्रमाने तुलना केली असता फक्त पर्याय ${correct_option_letter}$ "
        f"मधील संख्यांशी या संख्या तंतोतंत जुळतात.<br>"
        f"इतर पर्यायातील किमान एक तरी संख्या वेगळी आहे.<br>"
        f"म्हणून पर्याय ${correct_option_letter}$ हे उत्तर <br>"
    )

    return eng_sol + "\n" + mar_sol


def main():
    filename = "Intern_09030201_102_1_Assign2_Siddhi_Manjarekar.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = ["Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
              "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
              "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
              "Difficulty Level", "Question (Image/ Audio/ Video)", "Contributor's Registered mailId",
              "Solution (Text Only)", "Solution (Image/ Audio/ Video)", "Variation Number"]
    sheet.append(header)

    try:
        num_questions = int(input("How many questions do you want to enter? "))
    except ValueError:
        print("Invalid input. Please enter a number.")
        return

    generated_questions = set()
    i = 1

    while i <= num_questions:
        days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        rainfall = [random.randint(0, 50) for _ in range(7)]

        question_key = tuple(rainfall)
        if question_key in generated_questions:
            continue
        generated_questions.add(question_key)

        chart_base64 = create_chart_base64(rainfall, days, i)
        
        if not chart_base64:
            print(f"  Warning: Failed to generate chart for question {i}")
            continue

        scale = 5
        eng_question = "Daily rainfall is measured for a week from Monday to Sunday and the bar graph is plotted as shown below. From this bar graph decide the daily quantities of rainfall?"
        mar_question = "एका सप्ताहातील रोज पडणाऱ्या पावसाची मोजणी सोमवार ते रविवार अशी केली आणि त्या माहिती साठी खाली दाखविलेला स्तंभालेख तयार केला. या स्तंभालेखावरून रोज पडलेला पाऊस किती होता?"
        
        eng_title = "Daily Rainfall"
        mar_title = normalize_marathi_text("दैनंदिन पाऊस")
        
        eng_scale = f"Scale: 1 unit = {scale} mm"
        mar_scale = normalize_marathi_text(f"प्रमाण: 1 एकक = {scale} मिमी")
        
        eng_xlabel = "Day of the Week"
        mar_xlabel = normalize_marathi_text("आठवड्याचे दिवस")
        
        eng_ylabel = "Rainfall"
        mar_ylabel = normalize_marathi_text("पाऊस")
        
        eng_days_spaced = r"$\hspace{0.8cm}$ Mon $\hspace{0.8cm}$ Tue $\hspace{0.8cm}$ Wed $\hspace{0.8cm}$ Thu $\hspace{0.8cm}$ Fri $\hspace{0.8cm}$ Sat $\hspace{0.8cm}$ Sun"
        mar_days_spaced = r"$\hspace{0.8cm}$ सोम $\hspace{0.8cm}$ मंगळ $\hspace{0.8cm}$ बुध $\hspace{0.8cm}$ गुरु $\hspace{0.8cm}$ शुक्र $\hspace{0.8cm}$ शनि $\hspace{0.8cm}$ रवि"

        question_text = (
            f'<p>{eng_question}#{normalize_marathi_text(mar_question)}</p>'
            f'<div style="text-align:center;border:1px solid #ccc;padding:10px;max-width:550px;">'
            f'<span style="font-size:20px;">{eng_title} / {mar_title}</span><br>'
            f'<span style="font-size:15px;">({eng_scale} / {mar_scale})</span>'
            f'<div style="display:flex;align-items:center;justify-content:center;margin-top:5px;gap:0;">'
            f'<div style="writing-mode:vertical-rl;transform:rotate(180deg);'
            f'font-size:15px;white-space:nowrap;margin:0;padding:0;line-height:1;">'
            f'{eng_ylabel} (mm) / {mar_ylabel} (मिमी)</div>'
            f'<img src="{chart_base64}" alt="Chart" width="450" '
            f'style="display:block;margin:0;padding:0;">'
            f'</div>'
            f'<p style="font-size:12px;margin:8px 0 3px;font-weight:bold;">'
            f'{eng_days_spaced}<br>{normalize_marathi_text(mar_days_spaced)}</p>'
            f'<p style="font-size:15px;margin-top:8px;">{eng_xlabel} / {mar_xlabel}</p>'
            f'</div>'
        )

        correct_ans_values = f"${',\\\\ '.join(map(str, rainfall))}$"

        wrong_answers_list = generate_wrong_answers(rainfall, correct_ans_values)
        if not wrong_answers_list:
            continue

        options = [correct_ans_values] + wrong_answers_list
        option_letters = ['A', 'B', 'C', 'D']
        random.shuffle(options)
        correct_option_letter = option_letters[options.index(correct_ans_values)]

        formatted_options = [f"Option ${letter}$.<br>#पर्याय ${letter}$.<br>{value} mm<br>"
                             for letter, value in zip(option_letters, options)]

        correct_answer_formatted = formatted_options[options.index(correct_ans_values)]
        wrong_answers_formatted = [opt for opt in formatted_options if opt != correct_answer_formatted]

        solution = create_solution_text(rainfall, correct_option_letter)

        row_data = [i, "Text", 1, "09030201", question_text, correct_answer_formatted,
                    "", "", "", wrong_answers_formatted[0], wrong_answers_formatted[1], 
                    wrong_answers_formatted[2], 300, 1, "", "siddhimanjarekar274@gmail.com",
                    solution, "", 102]
        sheet.append(row_data)

        i += 1

    sheet.append(["****"])
    
    workbook.save(filename)
    print(f"\n✓✓✓ Successfully generated {num_questions} questions in '{filename}'!")


if __name__ == "__main__":
    main()