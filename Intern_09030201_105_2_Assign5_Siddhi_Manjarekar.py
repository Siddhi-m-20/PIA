import random
import unicodedata
import openpyxl
import warnings
import os

warnings.filterwarnings('ignore')

def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)


def create_pictograph_table(quantities, stations_eng, stations_mar, scale=40):
    
    face_image = '<img src="https://portal.coepvlab.ac.in/VirtualMathsLab/media/090201_108_emoji1.png" style="width:24px;height:24px;vertical-align:middle;">'
    
    table_html = '<table border="1" style="border-collapse:collapse;width:100%;text-align:left;">'
    
    table_html += '<tr style="background-color:#f0f0f0;">'
    table_html += '<th style="padding:8px;font-size:14px;">Station Name / स्थानकाचे नाव</th>'
    table_html += '<th style="padding:8px;font-size:14px;">Number of Travelers / प्रवाशांची संख्या</th>'
    table_html += '</tr>'
    
    for station_eng, station_mar, qty in zip(stations_eng, stations_mar, quantities):
        table_html += '<tr>'
        table_html += f'<td style="padding:8px;font-size:14px;">{station_eng} / {station_mar}</td>'
        table_html += f'<td style="padding:8px;">{face_image * qty}</td>'
        table_html += '</tr>'
    
    table_html += '</table>'
    return table_html


def create_chart_svg(quantities, stations_eng, stations_mar, scale=40):
    y_max = max(400, int(max([q * scale for q in quantities]) + 40))
    chart_height = 280
    chart_width = 500
    
    actual_quantities = [q * scale for q in quantities]
    bar_heights = [(q / y_max) * chart_height for q in actual_quantities]
    
    bar_colors = ['&num;4472C4', '&num;ED7D31', '&num;A5A5A5', '&num;FFC000', '&num;70AD47']
    
    bar_width = 60
    bar_spacing = 90
    start_x = 95
    
    svg_height = chart_height + 100  
    
    svg = f'<svg width=\'550\' height=\'{svg_height}\' xmlns=\'http://www.w3.org/2000/svg\'>'
    svg += '<rect width=\'100%\' height=\'100%\' fill=\'white\'/>'
    
    for i in range(0, y_max + 1, 40):
        y_pos = chart_height - (i / y_max) * chart_height + 20
        if i % 80 == 0:
            svg += f'<line x1=\'60\' y1=\'{y_pos}\' x2=\'530\' y2=\'{y_pos}\' stroke=\'gray\' stroke-width=\'1\' opacity=\'0.8\'/>'
            svg += f'<text x=\'50\' y=\'{y_pos + 4}\' text-anchor=\'end\' font-size=\'13\' font-family=\'Arial\'>{i}</text>'
        else:
            svg += f'<line x1=\'60\' y1=\'{y_pos}\' x2=\'530\' y2=\'{y_pos}\' stroke=\'gray\' stroke-width=\'0.5\' opacity=\'0.6\'/>'
    
    svg += f'<line x1=\'60\' y1=\'20\' x2=\'60\' y2=\'{chart_height + 20}\' stroke=\'black\' stroke-width=\'1\'/>'
    svg += f'<line x1=\'60\' y1=\'{chart_height + 20}\' x2=\'530\' y2=\'{chart_height + 20}\' stroke=\'black\' stroke-width=\'1\'/>'
    
    for i, (height, color, qty) in enumerate(zip(bar_heights, bar_colors, actual_quantities)):
        x_pos = start_x + i * bar_spacing
        bar_y = chart_height + 20 - height
        
        svg += f'<rect x=\'{x_pos}\' y=\'{bar_y}\' width=\'{bar_width}\' height=\'{height}\' fill=\'{color}\' stroke=\'black\' stroke-width=\'0.7\'/>'
    
    label_y_eng = chart_height + 40
    for i, station_eng in enumerate(stations_eng):
        x_center = start_x + i * bar_spacing + (bar_width / 2)
        svg += f'<text x=\'{x_center}\' y=\'{label_y_eng}\' text-anchor=\'middle\' font-size=\'12\' font-family=\'Arial\' font-weight=\'bold\'>{station_eng}</text>'
    
    label_y_mar = chart_height + 58
    for i, station_mar in enumerate(stations_mar):
        x_center = start_x + i * bar_spacing + (bar_width / 2)
        svg += f'<text x=\'{x_center}\' y=\'{label_y_mar}\' text-anchor=\'middle\' font-size=\'12\' font-family=\'Arial\' font-weight=\'bold\'>{normalize_marathi_text(station_mar)}</text>'
    
    svg += '</svg>'
    
    return svg


def generate_wrong_answers(quantities, stations_eng, stations_mar, scale=40, all_stations_eng=None, all_stations_mar=None):
    
    wrong_answers_set = set()
    wrong_options = []
    attempts = 0
    max_attempts = 100
    
    if all_stations_eng is None:
        all_stations_eng = stations_eng
    if all_stations_mar is None:
        all_stations_mar = stations_mar
    
    alt_stations_eng = [s for s in all_stations_eng if s not in stations_eng]
    alt_stations_mar = [s for s in all_stations_mar if s not in stations_mar]
    
    alt_scales = [20, 30, 50, 60]
    
    while len(wrong_options) < 3 and attempts < max_attempts:
        attempts += 1
        temp_quantities = quantities[:]
        temp_stations_eng = stations_eng[:]
        temp_stations_mar = stations_mar[:]
        temp_scale = scale
        
        method = random.randint(0, 6)

        if method == 0:
            idx = random.randrange(len(temp_quantities))
            change = random.choice([-3, -2, -1, 1, 2, 3])
            temp_quantities[idx] = max(1, min(10, temp_quantities[idx] + change))
        elif method == 1:
            random.shuffle(temp_quantities)
        elif method == 2:
            idx1, idx2 = random.sample(range(len(temp_quantities)), 2)
            temp_quantities[idx1], temp_quantities[idx2] = temp_quantities[idx2], temp_quantities[idx1]
        elif method == 3:
            indices = random.sample(range(len(temp_quantities)), 2)
            for idx in indices:
                change = random.choice([-2, -1, 1, 2])
                temp_quantities[idx] = max(1, min(10, temp_quantities[idx] + change))
        elif method == 4:
            if alt_stations_eng and alt_stations_mar:
                num_changes = random.randint(1, min(3, len(alt_stations_eng)))
                indices_to_change = random.sample(range(5), num_changes)
                available_eng = alt_stations_eng[:]
                available_mar = alt_stations_mar[:]
                
                for idx in indices_to_change:
                    if available_eng and available_mar:
                        new_eng = random.choice(available_eng)
                        eng_idx = all_stations_eng.index(new_eng)
                        new_mar = all_stations_mar[eng_idx]
                        temp_stations_eng[idx] = new_eng
                        temp_stations_mar[idx] = new_mar
                        available_eng.remove(new_eng)
                        available_mar.remove(new_mar)
        elif method == 5:
            temp_scale = random.choice(alt_scales)
        else:
            idx = random.randrange(len(temp_quantities))
            change = random.choice([-2, -1, 1, 2])
            temp_quantities[idx] = max(1, min(10, temp_quantities[idx] + change))
            
            if alt_stations_eng and alt_stations_mar:
                change_idx = random.randint(0, 4)
                available = [s for s in alt_stations_eng if s not in temp_stations_eng]
                if available:
                    new_eng = random.choice(available)
                    eng_idx = all_stations_eng.index(new_eng)
                    new_mar = all_stations_mar[eng_idx]
                    temp_stations_eng[change_idx] = new_eng
                    temp_stations_mar[change_idx] = new_mar

        wrong_key = (tuple(temp_quantities), tuple(temp_stations_eng), temp_scale)
        correct_key = (tuple(quantities), tuple(stations_eng), scale)
        
        if wrong_key != correct_key and wrong_key not in wrong_answers_set:
            wrong_answers_set.add(wrong_key)
            wrong_options.append({
                'quantities': temp_quantities[:],
                'stations_eng': temp_stations_eng[:],
                'stations_mar': temp_stations_mar[:],
                'scale': temp_scale
            })

    return wrong_options if len(wrong_options) == 3 else None


def create_solution_text(quantities, stations_eng, stations_mar, correct_option_letter, scale=40):
    q0_val = quantities[0] * scale
    q1_val = quantities[1] * scale
    q2_val = quantities[2] * scale
    q3_val = quantities[3] * scale
    q4_val = quantities[4] * scale
    
    face_symbol = '<img src="https://portal.coepvlab.ac.in/VirtualMathsLab/media/090201_108_emoji1.png" style="width:20px;height:20px;vertical-align:middle;">'
    
    eng_sol = (
        f"Ans : ${correct_option_letter}$.<br>"
        f"From the given pictograph, we get following information.<br>"
        f"$1)$ It is about the travelers by train boarding from $5$ different stations.<br>"
        f"$2)$ Scale used to represent the number of travelers is $1$ symbol of face {face_symbol} $= {scale}$ passengers.<br>"
        f"$3)$ Quantity of travelers boarding at different stations is as follows :<br>"
        f"{stations_eng[0]} $:$ ${quantities[0]}$ symbols of face $:\\ {quantities[0]} \\times {scale}= {q0_val}$ travelers,<br>"
        f"{stations_eng[1]} $:$ ${quantities[1]}$ symbols of face $:\\ {quantities[1]} \\times {scale}= {q1_val}$ travelers,<br>"
        f"{stations_eng[2]} $:$ ${quantities[2]}$ symbols of face $:\\ {quantities[2]} \\times {scale}= {q2_val}$ travelers,<br>"
        f"{stations_eng[3]} $:$ ${quantities[3]}$ symbols of face $:\\ {quantities[3]} \\times {scale}= {q3_val}$ travelers and,<br>"
        f"{stations_eng[4]} $:$ ${quantities[4]}$ symbols of face $:\\ {quantities[4]} \\times {scale}={q4_val}$ travelers.<br>"
        f"For deciding the correct bar graph representing this information correctly, we need to check the given bar graphs in the options for following conditions, based on above information.<br>"
        f"$i)$ It has $5$ bars corresponding to $5$ stations mentioned in given pictograph.<br>"
        f"$ii)$ All bars are equally spaced and of uniform width.<br>"
        f"$iii)$ Based on the scale mentioned in bar graph and the height of each bar must result into the given number of travelers boarding at each station.<br>"
        f"$iv)$ Scale is mentioned in the graph.<br>"
        f"By checking each bar graph in options for compliance of conditions $(i)$ to $(iv)$ mentioned above, we can conclude that, only bar graph given in option ${correct_option_letter}$, satisfies these conditions.<br>"
        f"In all other bar graphs, either the boarding station names don't match, or number of stations differs or the quantities of travelers differ or the scale used differs.<br>"
        f"Hence, option ${correct_option_letter}$ is the answer.<br>#"
    )

    mar_sol = (
        f"उत्तर : ${correct_option_letter}$.<br>"
        f"दिलेल्या चित्रालेखा वरून आपल्याला पुढील माहिती मिळते.<br>"
        f"$1)$ $5$ वेगवेगळ्या स्थानका वरून रेल्वे मधून प्रवास करणाऱ्या प्रवाशांची माहिती दिली आहे.<br>"
        f"$2)$ या चित्रालेखात वापरलेले प्रमाण  $1$ चेहऱ्याचे चिन्ह {face_symbol} $= {scale}$ प्रवासी.<br>"
        f"$3)$ प्रत्येक स्थानकावरून रेल्वे मध्ये येणाऱ्या प्रवाशांची संख्या पुढील प्रमाणे आहे :<br>"
        f"{stations_mar[0]} $: {quantities[0]}$ चेहेऱ्याची चिन्हे $:\\ {quantities[0]} \\times {scale}= {q0_val}$ प्रवासी,<br>"
        f"{stations_mar[1]} $: {quantities[1]}$ चेहेऱ्याची चिन्हे $:\\ {quantities[1]} \\times {scale}= {q1_val}$ प्रवासी,<br>"
        f"{stations_mar[2]} $: {quantities[2]}$ चेहेऱ्याची चिन्हे $:\\ {quantities[2]} \\times {scale}= {q2_val}$ प्रवासी,<br>"
        f"{stations_mar[3]} $: {quantities[3]}$ चेहेऱ्याची चिन्हे $:\\ {quantities[3]} \\times {scale}= {q3_val}$ प्रवासी आणि,<br>"
        f"{stations_mar[4]} $: {quantities[4]}$ चेहेऱ्याची चिन्हे $:\\ {quantities[4]} \\times {scale}={q4_val}$ प्रवासी.<br>"
        f"या माहित नुसार अचूक स्तंभालेख ठरविण्यासाठी आपल्याला पर्यायात दिलेले सर्व स्तंभालेख खालील बाबींसाठी तपासावे लागतील.<br>"
        f"$i)$ चित्रालेखात दिल्या नुसार त्यातील $5$ स्थानकांसाठी $5$ स्तंभ यात आहेत का.<br>"
        f"$ii)$ सर्व स्तंभांची रुंदी समान आहे का आणि कोणत्याही दोन लगतच्या स्तंभां मधील मोकळी जागा समान आहे का. <br>"
        f"$iii)$ स्तंभालेखात उल्लेख केलेल्या प्रमाणा नुसार त्यातील स्तंभांची उंची त्या त्या स्थानका नुसार योग्य प्रवासी संख्या दाखविते का. <br>"
        f"$iv)$ आलेखात प्रमाणाचा उल्लेख आहे.<br>"
        f"पर्यायात दिलेले सर्व स्तंभालेखा $(i)$ ते $(iv)$ या अटीं नुसार तपासले असता, आपण हे निश्चित पणे सांगू शकतो की, फक्त पर्याय ${correct_option_letter}$ मधील स्तंभालेख या अटी पूर्ण करतो.<br>"
        f"इतर सर्व स्तंभालेखात काही ठिकाणी स्थानकांची नावे जुळत नाहीत किंवा स्थानकांची संख्या वेगळी आहे किंवा प्रवासी संख्या जुळत नाही किंवा वापरलेले प्रमाण वेगळे आहे.<br>"
        f"म्हणून, पर्याय ${correct_option_letter}$ हे उत्तर.<br>"
    )

    return eng_sol + mar_sol


def create_option_chart_html(quantities, stations_eng, stations_mar, scale=40):
    chart_svg = create_chart_svg(quantities, stations_eng, stations_mar, scale)
    
    if not chart_svg:
        return ""
    
    eng_title = "Travelers Boarding Train"
    mar_title = normalize_marathi_text("रेल्वे प्रवासी")
    eng_scale = f"Scale: 1 unit = {scale} passengers"
    mar_scale = normalize_marathi_text(f"प्रमाण: 1 एकक = {scale} प्रवासी")
    eng_ylabel = "Y-Axis"
    mar_ylabel = normalize_marathi_text("Y-अक्ष")
    eng_xlabel = "X-Axis"
    mar_xlabel = normalize_marathi_text("X-अक्ष")
    
    complete_chart = (
        f'<div style=\'text-align:center;border:1px solid &num;ccc;padding:10px;max-width:650px;\'>'
        f'<span style=\'font-size:22px;\'>{eng_title} / {mar_title}</span><br>'
        f'<span style=\'font-size:15px;\'>({eng_scale} / {mar_scale})</span>'
        f'<div style=\'position:relative;margin-top:5px;\'>'
        f'<div style=\'position:absolute;top:-5px;left:50px;font-size:13px;\'>'
        f'{eng_ylabel} / {mar_ylabel}</div>'
        f'<div style=\'display:flex;align-items:center;justify-content:center;gap:0;\'>'
        f'<div style=\'writing-mode:vertical-rl;transform:rotate(180deg);'
        f'font-size:17px;white-space:nowrap;margin:0 5px 0 0;padding:0;line-height:1.2;'
        f'font-family:Arial, sans-serif;\'>'
        f'{normalize_marathi_text("Number of Travelers / प्रवाशांची संख्या")}</div>'
        f'{chart_svg}'
        f'</div>'
        f'<div style=\'position:absolute;bottom:0px;right:10px;font-size:13px;font-family:Arial;\'>'
        f'{eng_xlabel} / {mar_xlabel}</div>'
        f'</div>'
        f'<p style=\'font-size:17px;margin-top:8px;\'>Station Name / स्थानकाचे नाव</p>'
        f'</div>'
    )
    
    return complete_chart


def create_pictograph_html(quantities, stations_eng, stations_mar, scale=40):
    pictograph_table = create_pictograph_table(quantities, stations_eng, stations_mar, scale)
    
    face_image = '<img src="https://portal.coepvlab.ac.in/VirtualMathsLab/media/090201_108_emoji1.png" style="width:20px;height:20px;vertical-align:middle;">'
    
    eng_title = "Travelers Boarding / चढणारे प्रवासी"
    eng_scale = f"Scale: 1 {face_image} = {scale} passengers"
    mar_scale = normalize_marathi_text(f"प्रमाण: 1 {face_image} = {scale} प्रवासी")
    
    complete_pictograph = (
        f'<div style=\'text-align:center;border:1px solid &num;ccc;padding:10px;max-width:650px;\'>'
        f'<span style=\'font-size:22px;\'>{eng_title}</span><br>'
        f'<span style=\'font-size:15px;\'>({eng_scale} / {mar_scale})</span>'
        f'<div style=\'margin-top:10px;\'>'
        f'{pictograph_table}'
        f'</div>'
        f'</div>'
    )
    
    return complete_pictograph


def main():
    os.makedirs("09030201_105_2_Assign5_Siddhi_Manjarekar", exist_ok=True)
    filename = "09030201_105_2_Assign5_Siddhi_Manjarekar/Intern_09030201_105_2_Assign5_Siddhi_Manjarekar.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = ["Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
              "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
              "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
              "Difficulty Level", "Question (Image/ Audio/ Video)", "Contributor's Registered mailId",
              "Solution (Text Only)", "Solution (Image/ Audio/ Video)", "Variation Number"]
    sheet.append(header)

    try:
        num_questions = int(input("How many questions do you want to enter? "))
        if num_questions > 200:
            print("Warning: Requesting more than 200 questions may result in duplicates or take a long time.")
            proceed = input("Do you want to continue? (yes/no): ").strip().lower()
            if proceed != 'yes':
                return
    except ValueError:
        print("Invalid input. Please enter a number.")
        return

    all_stations_eng = ["Lonavala", "Kamshet", "Vadgaon", "Talegaon", "Chinchwad", 
                        "Pimpri", "Akurdi", "Kasarwadi", "Dapodi", "Dehuroad", 
                        "Bhegade wadi", "Ghorawadi", "Kaanhe", "Malawali"]
    all_stations_mar = ["लोणावळा", "कामशेत", "वडगाव", "तळेगाव", "चिंचवड",
                        "पिंपरी", "आकुर्डी", "कासारवाडी", "दापोडी", "देहूरोड",
                        "भेगडे वाडी", "घोरावाडी", "कान्हे", "मळवली"]
    
    option_letters = ['A', 'B', 'C', 'D']
    scale = 40
    
    generated_questions = set()
    i = 1
    rows_to_add = []
    max_retries = num_questions * 10
    retry_count = 0

    while i <= num_questions and retry_count < max_retries:
        retry_count += 1
        quantities = [random.randint(3, 10) for _ in range(5)]

        chosen_indices = random.sample(range(len(all_stations_eng)), 5)
        stations_eng5 = [all_stations_eng[idx] for idx in chosen_indices]
        stations_mar5 = [all_stations_mar[idx] for idx in chosen_indices]
        question_key = (tuple(quantities), tuple(stations_eng5))
        if question_key in generated_questions:
            continue
        
        wrong_answers_list = generate_wrong_answers(
            quantities, 
            stations_eng5, 
            stations_mar5, 
            scale,
            all_stations_eng,
            all_stations_mar
        )
        if not wrong_answers_list:
            continue
        
        generated_questions.add(question_key)

        pictograph_html = create_pictograph_html(quantities, stations_eng5, stations_mar5, scale)

        eng_question = (
            f"Here a pictograph is given. It's equivalent bar graph is to be drawn.<br>"
            f"Which of the following Bargraph is the correct representations for this?<br>"
        )

        mar_question = normalize_marathi_text(
            f"सोबत एक चित्रालेख दिलेला आहे. हाच चित्रालेख स्तंभलेखाच्या रूपात दाखवायचा आहे.<br>"
            f"तर पुढील पर्याया पैकी यातील कोणता स्तंभालेख बरोबर आहे?<br>"
        )

        question_text = f'<p>{eng_question}#{mar_question}{pictograph_html}</p>'

        correct_chart_html = create_option_chart_html(quantities, stations_eng5, stations_mar5, scale)
        if not correct_chart_html:
            continue
        
        wrong_charts_html = []
        for wrong_option in wrong_answers_list:
            wrong_chart_html = create_option_chart_html(
                wrong_option['quantities'], 
                wrong_option['stations_eng'], 
                wrong_option['stations_mar'], 
                wrong_option['scale']
            )
            if not wrong_chart_html:
                break
            wrong_charts_html.append(wrong_chart_html)
        
        if len(wrong_charts_html) != 3:
            continue

        all_charts = [correct_chart_html] + wrong_charts_html
        random.shuffle(all_charts)
        correct_option_letter = option_letters[all_charts.index(correct_chart_html)]

        formatted_options = [f"${letter}$.<br>{chart}"
                             for letter, chart in zip(option_letters, all_charts)]

        correct_answer_formatted = formatted_options[all_charts.index(correct_chart_html)]
        wrong_answers_formatted = [opt for opt in formatted_options if opt != correct_answer_formatted]

        solution = create_solution_text(quantities, stations_eng5, stations_mar5, correct_option_letter, scale)

        row_data = [i, "Text", 1, "09030201", question_text, correct_answer_formatted,
                    "", "", "", wrong_answers_formatted[0], wrong_answers_formatted[1], 
                    wrong_answers_formatted[2], 240, 2, "", "siddhimanjarekar2@gmail.com", 
                    solution, "", 105]
        
        rows_to_add.append(row_data)
        
        i += 1

    if i <= num_questions:
        print(f"\nWarning: Could only generate {i-1} unique questions out of {num_questions} requested.")
    
    for row in rows_to_add:
        sheet.append(row)
    
    sheet.append(["****"])
    
    workbook.save(filename)
    print(f"\nSuccessfully generated {len(rows_to_add)} questions in '{filename}'!")
    

if __name__ == "__main__":
    main()