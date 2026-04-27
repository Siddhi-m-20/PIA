import random
import unicodedata
import openpyxl
import warnings
import os

warnings.filterwarnings('ignore')

def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)


def create_chart_svg(quantities, grains_eng, grains_mar, scale=5):
    y_max = max(50, int(max(quantities) + 5))
    chart_height = 280
    chart_width = 400
    
    bar_heights = [(q / y_max) * chart_height for q in quantities]
    
    bar_colors = ['&num;4472C4', '&num;ED7D31', '&num;A5A5A5', '&num;FFC000']
    
    svg = f'<svg width=\'450\' height=\'320\' xmlns=\'http://www.w3.org/2000/svg\'>'
    svg += '<rect width=\'100%\' height=\'100%\' fill=\'white\'/>'
    
    for i in range(0, y_max + 1):
        y_pos = chart_height - (i / y_max) * chart_height + 20
        if i % 5 == 0:
            svg += f'<line x1=\'60\' y1=\'{y_pos}\' x2=\'430\' y2=\'{y_pos}\' stroke=\'gray\' stroke-width=\'1\' opacity=\'0.8\'/>'
            svg += f'<text x=\'50\' y=\'{y_pos + 4}\' text-anchor=\'end\' font-size=\'13\' font-family=\'Arial\'>{i}</text>'
        else:
            svg += f'<line x1=\'60\' y1=\'{y_pos}\' x2=\'430\' y2=\'{y_pos}\' stroke=\'gray\' stroke-width=\'0.5\' opacity=\'0.6\'/>'
    
    svg += f'<line x1=\'60\' y1=\'20\' x2=\'60\' y2=\'{chart_height + 20}\' stroke=\'black\' stroke-width=\'1\'/>'
    svg += f'<line x1=\'60\' y1=\'{chart_height + 20}\' x2=\'430\' y2=\'{chart_height + 20}\' stroke=\'black\' stroke-width=\'1\'/>'
    
    bar_width = 60
    bar_spacing = 90
    start_x = 95
    
    for i, (height, color, qty) in enumerate(zip(bar_heights, bar_colors, quantities)):
        x_pos = start_x + i * bar_spacing
        bar_y = chart_height + 20 - height
        
        svg += f'<rect x=\'{x_pos}\' y=\'{bar_y}\' width=\'{bar_width}\' height=\'{height}\' fill=\'{color}\' stroke=\'black\' stroke-width=\'0.7\'/>'
    
    svg += '</svg>'
    
    return svg


def generate_wrong_answers(quantities, max_attempts=50):
    """Generate 3 unique wrong answers that are different from correct answer"""
    wrong_answers_set = set()
    correct_tuple = tuple(quantities)
    attempts = 0
    
    while len(wrong_answers_set) < 3 and attempts < max_attempts:
        attempts += 1
        temp_quantities = quantities[:] 
        method = random.randint(0, 2)

        if method == 0:  
            idx = random.randrange(len(temp_quantities))
            change = random.choice([-11, -10, -9, -8, 8, 9, 10, 11])
            temp_quantities[idx] = max(10, min(60, temp_quantities[idx] + change))
        elif method == 1: 
            random.shuffle(temp_quantities)
        else:  
            idx1, idx2 = random.sample(range(len(temp_quantities)), 2)
            temp_quantities[idx1], temp_quantities[idx2] = temp_quantities[idx2], temp_quantities[idx1]

        wrong_ans_tuple = tuple(temp_quantities)
        
        if wrong_ans_tuple != correct_tuple:
            wrong_answers_set.add(wrong_ans_tuple)

    return [list(t) for t in wrong_answers_set] if len(wrong_answers_set) == 3 else None


def verify_answer_uniqueness(correct_answer, wrong_answers):
    """Verify all answers are unique - similar to Java validation"""
    all_answers = [tuple(correct_answer)] + [tuple(wa) for wa in wrong_answers]
    
    if len(set(all_answers)) != len(all_answers):
        return False
    
    correct_tuple = tuple(correct_answer)
    for wrong in wrong_answers:
        if correct_tuple == tuple(wrong):
            return False
    
    for i in range(len(wrong_answers)):
        for j in range(i + 1, len(wrong_answers)):
            if tuple(wrong_answers[i]) == tuple(wrong_answers[j]):
                return False
    
    return True


def create_solution_text(quantities, correct_option_letter, scale=5):
    q0_units = quantities[0]/scale
    q1_units = quantities[1]/scale
    q2_units = quantities[2]/scale
    q3_units = quantities[3]/scale
    
    eng_sol = (
        f"Ans : ${correct_option_letter}$.<br>"
        f"To decide which of the graphs is the correct option, we need to check correctness based on the following points.<br>"
        f"$1.$ All commodities and their names, for which data is given, must appear in the graph in the same sequence.<br>"
        f"$2.$ Based on the selected scale and the quantity of each commodity, height of the corresponding bar must match the data. This should be checked for all commodities.<br>"
        f"Scale selected for this bar graph is $1$ unit $= {scale}$ kg.<br>"
        f"Accordingly, corresponding to the quantity of each commodity sold, the height of the bar for it should be as follows:<br>"
        f"For Corn - ${quantities[0]}$ kg corresponds to ${q0_units:.1f}$ units,<br>"
        f"For Pea - ${quantities[1]}$ kg corresponds to ${q1_units:.1f}$ units,<br>"
        f"For Rice - ${quantities[2]}$ kg corresponds to ${q2_units:.1f}$ units, and<br>"
        f"For Sorghum - ${quantities[3]}$ kg corresponds to ${q3_units:.1f}$ units.<br>"
        f"By inspection, we can observe that the sequence as well as corresponding height for each commodity is matching only for option ${correct_option_letter}$.<br>"
        f"For all other options at least one of the parameter is not matching.<br>"
        f"$\\therefore$ ${correct_option_letter}$ is the answer.<br>#"
    )

    mar_sol = (
        f"उत्तर : ${correct_option_letter}$.<br>"
        f"योग्य पर्याय ठरवायला हवा असलेली दिलेली माहिती योग्य पद्धतीने दाखविते, हे सोपस्करणासाठी आपल्याला पुढील गोष्टी तपासाव्या लागतील.<br>"
        f"$1.$ दिलेल्या सर्व वस्तूंची त्यांच्या योग्य नावासह आकडेवारी त्याच क्रमाने स्तंभालेखात दाखविण्यात आलेली आहे.<br>"
        f"$2.$ स्तंभालेखासाठी ठरविण्यात आलेल्या प्रमाणा नुसार प्रत्येक वस्तूची किंमतीसाठी (संख्या) दाखविलेल्या स्तंभाची उंची योग्य (बरोबर)आहे.<br>"
        f"या स्तंभालेखासाठी घेतलेले प्रमाण $1$ एकक $= {scale}$ किग्रॅ.<br>"
        f"त्यानुसार विकलेल्या प्रत्येक धान्यासाठी दाखविलेल्या संबंधित उंची खालील प्रमाणे असायला हवी :<br>"
        f"मका - ${quantities[0]}$ किग्रॅ म्हणजे ${q0_units:.1f}$ एकक,<br>"
        f"वाटाणे - ${quantities[1]}$ किग्रॅ म्हणजे ${q1_units:.1f}$ एकक,<br>"
        f"तांदूळ - ${quantities[2]}$ किग्रॅ म्हणजे ${q2_units:.1f}$ एकक, आणि<br>"
        f"ज्वारी - ${quantities[3]}$ किग्रॅ म्हणजे ${q3_units:.1f}$ एकक.<br>"
        f"तपासणी नुसार आपल्याला हे लक्षात येते की सर्व वस्तूंचा क्रम दिल्यानुसार, आणि प्रत्येक स्तंभाची योग्य उंची फक्त पर्याय ${correct_option_letter}$ मधील स्तंभालेखा मध्येच आहे.<br>"
        f"इतर सर्व पर्यायात यातील आवश्यक अशी एक तरी गोष्ट नाही.<br>"
        f"$\\therefore$ ${correct_option_letter}$ हाच बरोबर पर्याय आहे, हे उत्तर.<br>"
    )

    return eng_sol + mar_sol

_MAR_TITLE = normalize_marathi_text("धान्य विक्री")

def create_option_chart_html(quantities, grains_eng, grains_mar, scale=5):
    """Create complete chart HTML with axis labels outside the SVG"""
    chart_svg = create_chart_svg(quantities, grains_eng, grains_mar, scale)
    
    if not chart_svg:
        return ""
    
    eng_title = "Grain Sales"
    eng_scale = f"Scale: 1 unit = {scale} kg"
    mar_scale = normalize_marathi_text(f"प्रमाण: 1 एकक = {scale} किग्रॅ")
    eng_ylabel = "Y-Axis"
    mar_ylabel = normalize_marathi_text("Y-अक्ष")
    eng_xlabel = "X-Axis"
    mar_xlabel = normalize_marathi_text("X-अक्ष")
    eng_grains_spaced = r"$\hspace{2.5cm}$ Corn $\hspace{1.5cm}$ Pea $\hspace{1.6cm}$ Rice $\hspace{1.5cm}$ Sorghum"
    mar_grains_spaced = normalize_marathi_text(r"$\hspace{2.3cm}$ मका $\hspace{1.5cm}$ वाटाणे $\hspace{1.3cm}$ तांदूळ $\hspace{1.5cm}$ ज्वारी")
    
    complete_chart = (
        f'<div style=\'text-align:center;border:1px solid &num;ccc;padding:10px;max-width:550px;\'>'
        f'<span style=\'font-size:22px;\'>{eng_title} / {_MAR_TITLE}</span><br>'
        f'<span style=\'font-size:15px;\'>({eng_scale} / {mar_scale})</span>'
        f'<div style=\'position:relative;margin-top:5px;\'>'
        f'<div style=\'position:absolute;top:-5px;left:50px;font-size:13px;\'>'
        f'{eng_ylabel} / {mar_ylabel}</div>'
        f'<div style=\'display:flex;align-items:center;justify-content:center;gap:0;\'>'
        f'<div style=\'writing-mode:vertical-rl;transform:rotate(180deg);'
        f'font-size:17px;white-space:nowrap;margin:0 5px 0 0;padding:0;line-height:1.2;'
        f'font-family:Arial, sans-serif;\'>'
        f'{normalize_marathi_text("Quantity (kg) / प्रमाण (किग्रॅ)")}</div>'
        f'{chart_svg}'
        f'</div>'
        f'<div style=\'position:absolute;bottom:0px;right:10px;font-size:13px;font-family:Arial;\'>'
        f'{eng_xlabel} / {mar_xlabel}</div>'
        f'</div>'
        f'<p style=\'font-size:12px;margin:8px 0 3px;font-weight:bold;\'>'
        f'{eng_grains_spaced}<br>{mar_grains_spaced}</p>'
        f'<p style=\'font-size:17px;margin-top:8px;\'>Type of Grain / धान्याचे प्रकार</p>'
        f'</div>'
    )
    
    return complete_chart


def main():
    os.makedirs("09030201_103_2_Assign3_Siddhi_Manjarekar", exist_ok=True)
    filename = "09030201_103_2_Assign3_Siddhi_Manjarekar/Intern_09030201_103_2_Assign3_Siddhi_Manjarekar.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = ["Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
              "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
              "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
              "Difficulty Level", "Question (Image/ Audio/ Video)", "Contributor's Registered mailId",
              "Solution (Text Only)", "Solution (Image/ Audio/ Video)", "Variation Number"]
    sheet.append(header)

    try:
        num_questions = int(input("How many questions do you want to enter? "))
    except ValueError:
        print("Invalid input. Please enter a number.")
        return

    grains_eng = ["Corn", "Pea", "Rice", "Sorghum"]
    grains_mar = ["मका", "वाटाणे", "तांदूळ", "ज्वारी"]
    option_letters = ['A', 'B', 'C', 'D']
    scale = 5
    
    generated_questions = {}
    i = 1
    rows_to_add = []
    max_retries = 100  

    while i <= num_questions:
        retry_count = 0
        question_generated = False
        
        while retry_count < max_retries and not question_generated:
            retry_count += 1
            
            quantities = [random.randint(10, 60) for _ in range(4)]

            question_key = tuple(quantities)
            mapsize_before = len(generated_questions)
            
            if question_key not in generated_questions:
                generated_questions[question_key] = i
                mapsize_after = len(generated_questions)
                
                if mapsize_before == mapsize_after:
                    print(f"Duplicate question detected for question {i}: {quantities}")
                    continue
            else:
                print(f"Duplicate question detected for question {i}: {quantities}")
                continue

            eng_question = (
                f"In a grocery shop, different grains sold in a day (in kilograms) are to be displayed using a bar graph. "
                f"The sales data is as follows. Which of the following bar graphs correctly represents the data?<br>"
                f"Corn - ${quantities[0]}$ kg, Pea - ${quantities[1]}$ kg, Rice - ${quantities[2]}$ kg, Sorghum - ${quantities[3]}$ kg<br>"
            )

            mar_question = (
                f"किराणा दुकानात एका दिवसात विकले जाणारे विविध धान्य किलोग्रॅममध्ये दाखवायचे आहेत. "
                f"विक्रीची माहिती खालीलप्रमाणे आहे. दिलेल्या पैकी कोणता चित्रलेख ही माहिती योग्य रीतीने दर्शवितो?<br>"
                f"मका - ${quantities[0]}$ किग्रॅ, वाटाणे - ${quantities[1]}$ किग्रॅ, तांदूळ - ${quantities[2]}$ किग्रॅ, ज्वारी - ${quantities[3]}$ किग्रॅ<br>"
            )

            question_text = f'<p>{eng_question}#{normalize_marathi_text(mar_question)}</p>'

            wrong_answers_list = generate_wrong_answers(quantities)
            if not wrong_answers_list:
                print(f"Failed to generate wrong answers for question {i}, retrying...")
                
                del generated_questions[question_key]
                continue

            if not verify_answer_uniqueness(quantities, wrong_answers_list):
                print(f"Duplicate answers detected for question {i}, retrying...")
                del generated_questions[question_key]
                continue

            correct_chart_html = create_option_chart_html(quantities, grains_eng, grains_mar, scale)
            if not correct_chart_html:
                print(f"Failed to create correct chart for question {i}, retrying...")
                del generated_questions[question_key]
                continue
            
            wrong_charts_html = []
            for wrong_quantities in wrong_answers_list:
                wrong_chart_html = create_option_chart_html(wrong_quantities, grains_eng, grains_mar, scale)
                if not wrong_chart_html:
                    break
                wrong_charts_html.append(wrong_chart_html)
            
            if len(wrong_charts_html) != 3:
                print(f"Failed to create all wrong charts for question {i}, retrying...")
                del generated_questions[question_key]
                continue

            all_charts = [correct_chart_html] + wrong_charts_html
            random.shuffle(all_charts)
            correct_option_letter = option_letters[all_charts.index(correct_chart_html)]

            formatted_options = [f"${letter}$<br>{chart}"
                                 for letter, chart in zip(option_letters, all_charts)]

            correct_answer_formatted = formatted_options[all_charts.index(correct_chart_html)]
            wrong_answers_formatted = [opt for opt in formatted_options if opt != correct_answer_formatted]

            solution = create_solution_text(quantities, correct_option_letter, scale)

            row_data = [i, "Text", 1, "09030201", question_text, correct_answer_formatted,
                        "", "", "", wrong_answers_formatted[0], wrong_answers_formatted[1], 
                        wrong_answers_formatted[2], 120, 2, "", "siddhimanjarekar274@gmail.com",
                        solution, "", 103]
            
            rows_to_add.append(row_data)
            question_generated = True
            
        if not question_generated:
            print(f"Warning: Could not generate question {i} after {max_retries} retries. Skipping.")
            if question_key in generated_questions:
                del generated_questions[question_key]
    
        i += 1

    for row in rows_to_add:
        sheet.append(row)
    
    sheet.append(["****"])
    
    workbook.save(filename)
    print(f"\nSuccessfully generated {len(rows_to_add)} questions in '{filename}'!")
    

if __name__ == "__main__":
    main()