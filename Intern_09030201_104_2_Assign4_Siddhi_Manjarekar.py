import random
import unicodedata
import openpyxl
import warnings
import os

warnings.filterwarnings('ignore')


def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)


def get_divisors(n):
    divisors = []
    for i in range(1, int(n ** 0.5) + 1):
        if n % i == 0:
            divisors.append(i)
            if i != n // i:
                divisors.append(n // i)
    return sorted(divisors)


def generate_values_from_scale(scale, num_values=4, min_height=2, max_height=15):
    multipliers = random.sample(range(min_height, max_height + 1), num_values)
    values = sorted([scale * m for m in multipliers])
    return values, multipliers


def validate_question_data(scale, values, min_height=2, max_height=15):
    heights = [v // scale for v in values]
    
    if not all(v % scale == 0 for v in values):
        return False, None, None, None
    
    min_h, max_h = min(heights), max(heights)
    
    if min_h < min_height or max_h > max_height:
        return False, None, None, None
    
    gcd = values[0]
    for v in values[1:]:
        while v:
            gcd, v = v, gcd % v
    
    divisors = get_divisors(gcd)
    
    if len(divisors) < 4:
        return False, None, None, None
    
    return True, heights, min_h, max_h

def generate_wrong_scales(correct_scale, values, needed=3, min_height=2, max_height=15):
    max_value = max(values)
    min_value = min(values)

    gcd = values[0]
    for v in values[1:]:
        while v:
            gcd, v = v, gcd % v

    all_divisors = sorted(get_divisors(gcd))
    if correct_scale in all_divisors:
        all_divisors.remove(correct_scale)

    candidate_typeA = []
    for d in all_divisors:
        min_h = min_value // d
        max_h = max_value // d
        if min_h < min_height or max_h > max_height:
            candidate_typeA.append(d)

    candidate_typeB = []
    for scale in range(2, max_value + 1):
        if scale == correct_scale or scale in all_divisors:
            continue
        num_divisible = sum(1 for v in values if v % scale == 0)
        if num_divisible >= 1 and num_divisible < len(values):
            min_h = min_value // scale
            max_h = max_value // scale
            if min_h >= 1 and max_h <= 20:
                candidate_typeB.append(scale)

    candidate_typeC = []
    common_scales = [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18, 20, 22, 24, 25, 30]

    for scale in common_scales:
        if scale == correct_scale or scale in all_divisors or scale > max_value:
            continue
        if not any(v % scale == 0 for v in values):
            min_h = min_value // scale
            max_h = max_value // scale
            if abs(scale - correct_scale) <= 4 or (min_h >= 1 and max_h <= 18):
                candidate_typeC.append(scale)

    for offset in [-3, -2, -1, 1, 2, 3, 4]:
        scale = correct_scale + offset
        if scale >= 2 and scale <= max_value and scale not in candidate_typeC:
            if scale != correct_scale and scale not in all_divisors:
                if not any(v % scale == 0 for v in values):
                    candidate_typeC.append(scale)

    wrong_scales = []

    if candidate_typeB:
        num_from_B = min(2, len(candidate_typeB))
        selected_B = random.sample(candidate_typeB, num_from_B)
        wrong_scales.extend(selected_B)

    if len(wrong_scales) < needed and candidate_typeA:
        pick = random.choice(candidate_typeA)
        if pick not in wrong_scales:
            wrong_scales.append(pick)

    if len(wrong_scales) < needed and candidate_typeC:
        remaining_needed = needed - len(wrong_scales)
        available_C = [c for c in candidate_typeC if c not in wrong_scales]
        num_from_C = min(remaining_needed, len(available_C))
        if num_from_C > 0:
            selected_C = random.sample(available_C, num_from_C)
            wrong_scales.extend(selected_C)

    if len(wrong_scales) < needed:
        combined_pool = []
        combined_pool.extend([d for d in candidate_typeA if d not in wrong_scales])
        combined_pool.extend([d for d in candidate_typeB if d not in wrong_scales])
        combined_pool.extend([d for d in candidate_typeC if d not in wrong_scales])
        combined_pool.extend([d for d in all_divisors if d not in wrong_scales and d not in combined_pool])

        while len(wrong_scales) < needed and combined_pool:
            pick = combined_pool.pop(random.randrange(len(combined_pool)))
            if pick not in wrong_scales:
                wrong_scales.append(pick)

    scale_cursor = 2
    while len(wrong_scales) < needed and scale_cursor <= max_value:
        if scale_cursor != correct_scale and scale_cursor not in wrong_scales and scale_cursor not in all_divisors:
            wrong_scales.append(scale_cursor)
        scale_cursor += 1

    wrong_scales = list(dict.fromkeys(wrong_scales))[:needed]
    random.shuffle(wrong_scales)

    return wrong_scales



def create_solution_text(values, correct_scale):
    max_value = max(values)
    min_value = min(values)
    
    gcd = values[0]
    for v in values[1:]:
        while v:
            gcd, v = v, gcd % v
    
    divisors = [d for d in get_divisors(gcd) if d <= max_value]
    test_scales = sorted(divisors)[:7]
    
    eng_calc_text, mar_calc_text = "", ""
    for scale in test_scales:
        min_h = min_value // scale
        max_h = max_value // scale
        eng_calc_text += f"For ${scale}$ student{'s' if scale > 1 else ''} $= 1$ unit, height of bars will vary from ${min_h}$ units to ${max_h}$ units<br>"
        mar_calc_text += f"${scale}$ विद्यार्थी $= 1$ एकक, या नुसार स्तंभांची उंची ${min_h}$ एकक ते ${max_h}$ एकक अशी बदलेल.<br>"
    
    correct_min_h = min_value // correct_scale
    correct_max_h = max_value // correct_scale
    values_str = ", ".join([str(v) for v in values])
    probable_scales_str = ", ".join([str(s) for s in divisors])
    
    eng_sol = (
        f"Ans: ${correct_scale}$ students $= 1$ unit.<br>"
        f"We know that the selected scale factor should be such that it can represent the corresponding values on the graph paper.<br>"
        f"Also, the selected scale should be such that all the <b>values will be completely divisible</b> by the scale selected.<br>"
        f"By observing the data values, we can test different scale factors:<br>"
        f"Data values for this graph are {values_str} students.<br>"
        f"For this set of data, only probable scale could be {probable_scales_str}.<br>"
        f"{eng_calc_text}"
        f"Excessively large as well as too small heights for the bars are to be avoided.<br>"
        f"So, we will select a scale of ${correct_scale}$ students $= 1$ unit.<br>"
        f"This will give us height of bars ranging from ${correct_min_h}$ units to ${correct_max_h}$ units.<br>"
        f"<b>Note: We will discard the values, where bar height becomes too large.</b><br>"
        f"Hence, the proper scale would be ${correct_scale}$ students $= 1$ unit.<br>#"
    )
    
    mar_sol = (
        f"उत्तर: ${correct_scale}$ विद्यार्थी $= 1$ एकक.<br>"
        f"आपल्याला हे माहिती आहे की, निवडलेले प्रमाण असे असावे, ज्यामुळे दिलेल्या किमतींसाठी तयार होणाऱ्या स्तंभांची उंची आलेख कागदावर व्यवस्थित मावेल.<br>"
        f"त्याच बरोबर निवडलेल्या प्रमाणाच्या संख्येने <b>दिलेल्या माहितीच्या संख्यांना पूर्ण भाग</b> जाईल आणि त्या स्तंभांची उंची आलेख कागदावर मावेल.<br>"
        f"दिलेल्या माहितीच्या संख्यांनुसार आपण विविध प्रमाण तपासू शकतो. जसे की:<br>"
        f"दिलेल्या माहितीच्या किमती नुसार विद्यार्थी संख्या {values_str} अशी आहे.<br>"
        f"या माहितीसाठी, फक्त {probable_scales_str} हे प्रमाण शक्य आहेत.<br>"
        f"{mar_calc_text}"
        f"प्रमाण निवडताना स्तंभाची उंची खूप मोठी किंवा खूप लहान असू नये.<br>"
        f"म्हणून आपण ${correct_scale}$ विद्यार्थी $= 1$ एकक असे प्रमाण निवडू.<br>"
        f"यामुळे स्तंभांची उंची ${correct_min_h}$ ते ${correct_max_h}$ एकक इतकीच असेल.<br>"
        f"<b>टीप - जे प्रमाण घेतल्याने स्तंभांची उंची खूप जास्त होईल, अशा प्रमाण किमती आपण विचारात घेणार नाही.</b><br>"
        f"म्हणून या स्तंभलेखासाठी प्रमाण ${correct_scale}$ विद्यार्थी $= 1$ एकक, हे उत्तर.<br>"
    )
    
    return eng_sol + mar_sol


def main():
    os.makedirs("09030201_104_2_Assign4_Siddhi_Manjarekar", exist_ok=True)
    filename = "09030201_104_2_Assign4_Siddhi_Manjarekar/Intern_09030201_104_2_Assign4_Siddhi_Manjarekar.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = ["Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
              "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
              "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
              "Difficulty Level", "Question (Image/ Audio/ Video)", "Contributor's Registered mailId",
              "Solution (Text Only)", "Solution (Image/ Audio/ Video)", "Variation Number"]
    sheet.append(header)

    try:
        num_questions = int(input("How many questions do you want to enter? "))
        if num_questions <= 0:
            print("Please enter a positive number.")
            return
    except ValueError:
        print("Invalid input. Please enter a number.")
        return

    schools_list = [
        ("Navnirman Vidyalay", "नवनिर्माण विद्यालयातील"),
        ("Shikshan Prasarak Vidyalay", "शिक्षण प्रसारक विद्यालयातील"),
        ("Bharatiya Vidya Bhavan", "भारतीय विद्या भवनातील"),
        ("Dnyandeep Vidyalay", "ज्ञानदीप विद्यालयातील"),
        ("Vidya Vikas Vidyalay", "विद्या विकास विद्यालयातील"),
        ("Saraswati Vidyamandir", "सरस्वती विद्यामंदिरातील"),
        ("Pragati Vidyalay", "प्रगती विद्यालयातील"),
        ("Ramkrishna Vidyalay", "रामकृष्ण विद्यालयातील"),
        ("Modern Vidyalay", "मॉडर्न विद्यालयातील"),
        ("Bal Bharati Vidyalay", "बाल भारती विद्यालयातील")
    ]

    games_list = [
        ("Chess", "बुद्धिबळ"),
        ("Badminton", "बॅडमिंटन"),
        ("Football", "फुटबॉल"),
        ("Hockey", "हॉकी"),
        ("Cricket", "क्रिकेट"),
        ("Basketball", "बास्केटबॉल"),
        ("Volleyball", "व्हॉलीबॉल"),
        ("Table Tennis", "टेबल टेनिस")
    ]

    possible_scales = [5, 6, 7, 8, 9, 10, 11, 12, 15, 16, 18, 20]
    
    used_combinations = set()
    i, attempts, max_attempts = 1, 0, num_questions * 1000

    while i <= num_questions and attempts < max_attempts:
        attempts += 1
        
        correct_scale = random.choice(possible_scales)
        values, heights = generate_values_from_scale(correct_scale, num_values=4)
        is_valid, bar_heights, min_h, max_h = validate_question_data(correct_scale, values)
        
        if not is_valid:
            continue
        
        combination_key = (correct_scale, tuple(values))
        if combination_key in used_combinations:
            continue
        
        wrong_scales = generate_wrong_scales(correct_scale, values, needed=3)
        if len(wrong_scales) < 3:
            continue
        
        used_combinations.add(combination_key)
        
        school_eng, school_mar = random.choice(schools_list)
        selected_games = random.sample(games_list, 4)

        game_value_pairs = list(zip(selected_games, values))
        random.shuffle(game_value_pairs)
        
        shuffled_games, shuffled_values = zip(*game_value_pairs)

        eng_data_lines = "<br>".join([f"{game[0]}: ${val}$ students" 
                                      for game, val in zip(shuffled_games, shuffled_values)])
        mar_data_lines = "<br>".join([f"{normalize_marathi_text(game[1])}: ${val}$ विद्यार्थी" 
                                      for game, val in zip(shuffled_games, shuffled_values)])

        eng_question = (
            f"Information is collected from {school_eng} about the game the students like most.<br>"
            f"The data collected is as follows:<br>"
            f"{eng_data_lines}<br>"
            f"A bar graph is to be drawn for this information.<br>"
            f"Decide the proper scale to be selected for this bar graph.<br>#"
        )

        mar_question = normalize_marathi_text(
            f"{school_mar} विद्यार्थ्यांकडून त्यांना आवडणाऱ्या खेळाची माहिती गोळा केली.<br>"
            f"ती माहिती खाली प्रमाणे:<br>"
            f"{mar_data_lines}<br>"
            f"ही माहिती दाखविण्यासाठी स्तंभलेखाचा वापर करायचा आहे.<br>"
            f"या स्तंभलेखासाठी वापरायचे योग्य प्रमाण कोणते असेल?<br>"
        )

        question_text = f"{eng_question}{mar_question}"
        correct_answer = f"${correct_scale}$ students $= 1$ unit<br>${correct_scale}$ विद्यार्थी $= 1$ एकक<br>"
        wrong_answers = [f"${scale}$ students $= 1$ unit<br>${scale}$ विद्यार्थी $= 1$ एकक<br>" 
                        for scale in wrong_scales]
        
        solution = create_solution_text(values, correct_scale)

        random.shuffle(wrong_answers)
        row_data = [i, "Text", 1, "09030201", question_text, correct_answer,
                    "", "", "", wrong_answers[0], wrong_answers[1], wrong_answers[2],
                    120, 2, "", "siddhimanjarekar2@gmail.com", solution, "", 104]

        sheet.append(row_data)
        i += 1

    sheet.append(["****"])
    workbook.save(filename)
    print(f"Successfully generated {i - 1} questions in '{filename}'!")


if __name__ == "__main__":
    main()