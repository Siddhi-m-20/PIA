import random
import unicodedata
import openpyxl
import warnings


warnings.filterwarnings('ignore')

SCENARIOS = [
    {
        "type": "Vehicles",
        "topic_eng": "the number of travellers travelling by different vehicles",
        "topic_mar": "किती प्रवासी कोणकोणत्या वाहनातून प्रवास करतात",
        "chart_title_eng": "Travellers by Vehicle", "chart_title_mar": "वाहनानुसार प्रवासी",
        "x_label_eng": "Vehicle", "x_label_mar": "वाहन",
        "y_label_eng": "Number of Travellers", "y_label_mar": "प्रवाशांची संख्या",
        "y_unit_eng": "travellers", "y_unit_mar": "प्रवासी",
        "items": [
            {"eng": "Bicycle", "mar": "सायकल"}, {"eng": "Bike", "mar": "दुचाकी"},
            {"eng": "Car", "mar": "कार"}, {"eng": "Bus", "mar": "बस"},
            {"eng": "Rickshaw", "mar": "रिक्षा"}, {"eng": "Truck", "mar": "ट्रक"},
            {"eng": "Train", "mar": "रेल्वे"}, {"eng": "Scooter", "mar": "स्कूटर"}
        ]
    },
    {
        "type": "Fruits",
        "topic_eng": "the quantity of fruit sales in a shop",
        "topic_mar": "दुकानात कोणत्या फळांची किती विक्री झाली",
        "chart_title_eng": "Fruit Sales", "chart_title_mar": "फळांची विक्री",
        "x_label_eng": "Fruit", "x_label_mar": "फळ",
        "y_label_eng": "Quantity", "y_label_mar": "प्रमाण",
        "y_unit_eng": "kg", "y_unit_mar": "किग्रॅ",
        "items": [
            {"eng": "Apple", "mar": "सफरचंद"}, {"eng": "Banana", "mar": "केळी"},
            {"eng": "Mango", "mar": "आंबा"}, {"eng": "Grapes", "mar": "द्राक्षे"},
            {"eng": "Orange", "mar": "संत्री"}, {"eng": "Papaya", "mar": "पपई"},
            {"eng": "Guava", "mar": "पेरू"}, {"eng": "Pineapple", "mar": "अननस"}
        ]
    },
    {
        "type": "Library",
        "topic_eng": "the number of books for different subjects in a library",
        "topic_mar": "ग्रंथालयात कोणत्या विषयाची किती पुस्तके आहेत",
        "chart_title_eng": "Books in Library", "chart_title_mar": "ग्रंथालयातील पुस्तके",
        "x_label_eng": "Subject", "x_label_mar": "विषय",
        "y_label_eng": "Number of Books", "y_label_mar": "पुस्तकांची संख्या",
        "y_unit_eng": "books", "y_unit_mar": "पुस्तके",
        "items": [
            {"eng": "History", "mar": "इतिहास"}, {"eng": "Science", "mar": "विज्ञान"},
            {"eng": "Math", "mar": "गणित"}, {"eng": "English", "mar": "इंग्रजी"},
            {"eng": "Marathi", "mar": "मराठी"}, {"eng": "Arts", "mar": "कला"},
            {"eng": "Geography", "mar": "भूगोल"}, {"eng": "Civics", "mar": "नागरिकशास्त्र"}
        ]
    },
    {
        "type": "Trees",
        "topic_eng": "the number of different trees planted",
        "topic_mar": "कोणत्या प्रकारची किती झाडे लावली",
        "chart_title_eng": "Tree Plantation", "chart_title_mar": "वृक्षारोपण",
        "x_label_eng": "Tree", "x_label_mar": "झाड",
        "y_label_eng": "Number of Trees", "y_label_mar": "झाडांची संख्या",
        "y_unit_eng": "trees", "y_unit_mar": "झाडे",
        "items": [
            {"eng": "Neem", "mar": "कडुनिंब"}, {"eng": "Banyan", "mar": "वड"},
            {"eng": "Peepal", "mar": "पिंपळ"}, {"eng": "Mango", "mar": "आंबा"},
            {"eng": "Tamarind", "mar": "चिंच"}, {"eng": "Gulmohar", "mar": "गुलमोहर"},
            {"eng": "Ashoka", "mar": "अशोक"}, {"eng": "Coconut", "mar": "नारळ"}
        ]
    }
]


def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)


def generate_dataset(scenario):
    num_items = random.randint(5, 6)
    selected_items = random.sample(scenario["items"], num_items)
    scale_options = [10, 20, 50, 100]
    scale = random.choice(scale_options)
    data = []
    used_values = set()
    for item in selected_items:
        attempts = 0
        while True:
            qty_units = random.randint(2, 9)
            qty = qty_units * scale
            if qty not in used_values or attempts > 5:
                used_values.add(qty)
                break
            attempts += 1
        data.append({
            "eng_name": item["eng"],
            "mar_name": item["mar"],
            "value": qty,
            "units": qty_units
        })
    random.shuffle(data)
    return data, scale


def create_bar_chart_svg(data, scale, scenario):
    chart_width = 700
    chart_height = 350
    top_margin = 60
    left_margin = 100
    bottom_margin = 70
    num_bars = len(data)
    available_width = chart_width - left_margin - 50
    bar_width = 50
    gap = (available_width - (num_bars * bar_width)) / (num_bars + 1)
    max_units = max(d['units'] for d in data)
    y_axis_max = max_units + 1
    graph_height = chart_height - top_margin - bottom_margin
    unit_pixel_height = graph_height / y_axis_max

    svg = f"<svg width='{chart_width}' height='{chart_height}' xmlns='http://www.w3.org/2000/svg' style='background-color:white; border:1px solid &num;ddd;'>"
    svg += "<rect width='100%' height='100%' fill='white'/>"

    for i in range(y_axis_max + 1):
        y_pos = chart_height - bottom_margin - (i * unit_pixel_height)
        label_val = i * scale
        svg += f"<line x1='{left_margin}' y1='{y_pos}' x2='{chart_width - 20}' y2='{y_pos}' stroke='&num;e0e0e0' stroke-width='1' />"
        svg += f"<text x='{left_margin - 10}' y='{y_pos + 5}' text-anchor='end' font-family='Arial' font-size='12'>{label_val}</text>"

    svg += f"<line x1='{left_margin}' y1='{top_margin}' x2='{left_margin}' y2='{chart_height - bottom_margin}' stroke='black' stroke-width='2'/>"
    svg += f"<line x1='{left_margin}' y1='{chart_height - bottom_margin}' x2='{chart_width - 20}' y2='{chart_height - bottom_margin}' stroke='black' stroke-width='2'/>"

    mid_y = (chart_height / 2)
    y_label = f"{scenario['y_label_eng']} / {scenario['y_label_mar']}"
    svg += f"<text transform='rotate(-90, 30, {mid_y})' x='30' y='{mid_y}' text-anchor='middle' font-family='Arial' font-size='14' font-weight='bold'>{y_label}</text>"

    mid_x = left_margin + (available_width / 2)
    x_label = f"{scenario['x_label_eng']} / {scenario['x_label_mar']}"
    svg += f"<text x='{mid_x}' y='{chart_height - 15}' text-anchor='middle' font-family='Arial' font-size='14' font-weight='bold'>{x_label}</text>"

    bar_colors = ["SteelBlue", "Peru", "IndianRed", "CadetBlue", "SeaGreen", "GoldenRod", "SlateBlue", "DarkKhaki"]

    current_x = left_margin + gap
    for i, item in enumerate(data):
        bar_h = item['units'] * unit_pixel_height
        y_pos = chart_height - bottom_margin - bar_h
        color = bar_colors[i % len(bar_colors)]
        svg += f"<rect x='{current_x}' y='{y_pos}' width='{bar_width}' height='{bar_h}' fill='{color}' stroke='black' stroke-width='1'/>"
        svg += f"<text x='{current_x + bar_width / 2}' y='{y_pos - 5}' text-anchor='middle' font-family='Arial' font-size='12' font-weight='bold'>{item['value']}</text>"
        svg += f"<text x='{current_x + bar_width / 2}' y='{chart_height - bottom_margin + 20}' text-anchor='middle' font-family='Arial' font-size='11'>{item['eng_name']}</text>"
        svg += f"<text x='{current_x + bar_width / 2}' y='{chart_height - bottom_margin + 38}' text-anchor='middle' font-family='Arial' font-size='11'>{normalize_marathi_text(item['mar_name'])}</text>"
        current_x += bar_width + gap

    title = f"{scenario['chart_title_eng']} / {scenario['chart_title_mar']}"
    svg += f"<text x='{chart_width / 2}' y='30' text-anchor='middle' font-family='Arial' font-size='16' font-weight='bold'>{title}</text>"
    svg += f"<rect x='{chart_width - 180}' y='5' width='170' height='30' fill='white' stroke='black' stroke-width='1'/>"
    svg += f"<text x='{chart_width - 95}' y='25' text-anchor='middle' font-family='Arial' font-size='12'>Scale: 1 unit = {scale} {scenario['y_unit_eng']}</text>"
    svg += "</svg>"
    return svg


def generate_question_group(sr_no_start, variation_num):
    scenario = random.choice(SCENARIOS)
    data, scale = generate_dataset(scenario)
    svg_code = create_bar_chart_svg(data, scale, scenario)

    # Pre-calculations for questions
    sorted_asc = sorted(data, key=lambda x: x['value'])
    sorted_desc = sorted(data, key=lambda x: x['value'], reverse=True)

    max_item = sorted_desc[0]
    min_item = sorted_asc[0]
    second_max = sorted_desc[1]
    third_max = sorted_desc[2]
    second_min = sorted_asc[1]

    total_val = sum(x['value'] for x in data)
    diff_max_min = max_item['value'] - min_item['value']
    sum_max_min = max_item['value'] + min_item['value']
    sum_2nd_3rd_max = second_max['value'] + third_max['value']
    top_3_sum = sorted_desc[0]['value'] + sorted_desc[1]['value'] + sorted_desc[2]['value']

    rows = []

    main_q_text = f"A Bar Graph for {scenario['topic_eng']} is as shown in the following figure. With help of this graph answer the following questions<br>#सोबतच्या स्तंभालेखात {scenario['topic_mar']} हे दाखविले आहे. या स्तंभालेखाच्या आधारे खाली विचारलेल्या प्रश्नांची उत्तरे योग्य ठिकाणी खूण करून सांगा.<br>{svg_code}"

    main_sol_text = f"From the bar graph we get following information.<br>$1.$ Scale used for the graph.<br>$2.$ Different {scenario['x_label_eng']}s.<br>$3.$ {scenario['y_label_eng']} for different {scenario['x_label_eng']}s.<br>Based on this information, we will answer the given questions.<br>#उत्तर : <br>दिलेल्या स्तंभालेखा वरून आपल्याला पुढील माहिती मिळते.<br>$1.$ स्तंभालेखा साठी वापरलेले प्रमाण. <br>$2.$ वेगवेगळी {scenario['x_label_mar']}.<br>$3.$ प्रत्येक {scenario['x_label_mar']}साठी {scenario['y_label_mar']}.<br>या माहिती वरून आपण विचारलेल्या प्रश्नांची उत्तरे शोधू."

    main_row = [
        sr_no_start, "Text", 9, "09030201", main_q_text,
        "", "", "", "",
        "", "", "",
        240, 3, "", "siddhimanjarekar274@gmail.com",
        main_sol_text, "", variation_num
    ]
    rows.append(main_row)
    current_sr_no = sr_no_start + 1

    questions_pool = []

    q1 = {
        "q": f"This Bar Graph is for . . . . <br>#हा स्तंभालेख . . . . . साठी आहे.",
        "corr": f"showing {scenario['topic_eng']}",
        "wrong": [f"showing {scenario['x_label_eng']} names only", f"showing map of city", "showing price list"],
        "sol": f"Ans : showing {scenario['topic_eng']}.<br>As mentioned in the problem statement, the bar graph shows {scenario['topic_eng']}.<br>Hence, showing {scenario['topic_eng']}, is the answer.<br>#उत्तर : {scenario['topic_mar']} दाखविण्यासाठी.<br>प्रश्नात दिल्यानुसार {scenario['topic_mar']} दाखविण्यासाठी हे स्तंभालेख आहे.<br>म्हणून, {scenario['topic_mar']} दाखविण्यासाठी, हे उत्तर."
    }
    questions_pool.append(q1)

    w_scales = [scale * 2, scale * 10, int(scale / 2) if scale > 1 else 5]
    q2 = {
        "q": f"Scale used in this graph is . . . . <br>#या स्तंभालेखासाठी वापरलेले प्रमाण . . . . आहे.<br>$1$ unit = . . . {scenario['y_unit_eng']}.<br>#$1$ एकक = ... {scenario['y_unit_mar']} असे आहे.",
        "corr": f"$1$ unit = {scale} {scenario['y_unit_eng']}",
        "wrong": [f"$1$ unit = {w} {scenario['y_unit_eng']}" for w in w_scales],
        "sol": f"Ans : $1$ unit = {scale} {scenario['y_unit_eng']}.<br>As shown in bar graph, the scale used is, $1$ unit = {scale} {scenario['y_unit_eng']}.<br>Hence, $1$ unit = {scale} {scenario['y_unit_eng']}, is the answer.<br>#उत्तर : $1$ एकक = {scale} {scenario['y_unit_mar']}.<br>स्तंभलेखात दाखविल्या नुसार वापरलेले प्रमाण $1$ एकक = {scale} {scenario['y_unit_mar']} असे आहे.<br>म्हणून, $1$ एकक = {scale} {scenario['y_unit_mar']}, हे उत्तर."
    }
    questions_pool.append(q2)

    w_top3 = [top_3_sum + scale, top_3_sum - scale, top_3_sum + 10]
    q3 = {
        "q": f"How many are the total of {scenario['y_label_eng']} for top 3 values? <br>#सगळ्यात जास्त मूल्य असलेल्या पहिल्या तीन {scenario['x_label_mar']}ची एकूण {scenario['y_label_mar']} किती?",
        "corr": str(top_3_sum),
        "wrong": [str(w) for w in w_top3],
        "sol": f"Ans : {top_3_sum}.<br>From the given graph we can see that<br>Value for 1st max ({sorted_desc[0]['eng_name']}) $= {sorted_desc[0]['value']}$<br>Value for 2nd max ({sorted_desc[1]['eng_name']}) $= {sorted_desc[1]['value']}$<br>Value for 3rd max ({sorted_desc[2]['eng_name']}) $= {sorted_desc[2]['value']}$<br>By taking the total of these three numbers we get <br>${sorted_desc[0]['value']} + {sorted_desc[1]['value']} + {sorted_desc[2]['value']} = {top_3_sum}$.<br>Hence, {top_3_sum} is the answer.<br>#उत्तर : {top_3_sum}.<br>दिलेल्या स्तंभालेखानुसार <br>सगळ्यात जास्त मूल्य असलेल्या पहिल्या {scenario['x_label_mar']}ची ( {normalize_marathi_text(sorted_desc[0]['mar_name'])} ) संख्या $= {sorted_desc[0]['value']}$<br>दुसऱ्या {scenario['x_label_mar']}ची ( {normalize_marathi_text(sorted_desc[1]['mar_name'])} ) संख्या $= {sorted_desc[1]['value']}$<br>तिसऱ्या {scenario['x_label_mar']}ची ( {normalize_marathi_text(sorted_desc[2]['mar_name'])} ) संख्या $= {sorted_desc[2]['value']}$<br>या तीनही संख्येची बेरीज करून <br>${sorted_desc[0]['value']} + {sorted_desc[1]['value']} + {sorted_desc[2]['value']} = {top_3_sum}$ असे उत्तर मिळते."
    }
    questions_pool.append(q3)

    wrong_names = [x['eng_name'] for x in data if x != min_item]
    random.shuffle(wrong_names)
    q4 = {
        "q": f"Which {scenario['x_label_eng']} has minimum {scenario['y_label_eng']}?<br>#{scenario['y_label_mar']} सर्वात कमी कोणत्या {scenario['x_label_mar']}चे आहे?",
        "corr": min_item['eng_name'],
        "wrong": wrong_names[:3],
        "sol": f"Ans : {min_item['eng_name']}.<br>From the given graph we can see that<br>{min_item['value']} is the value for {min_item['eng_name']} which is least.<br>Hence {min_item['eng_name']} is the answer.<br>#उत्तर : {normalize_marathi_text(min_item['mar_name'])}.<br>दिलेल्या स्तंभालेखानुसार <br>{min_item['value']} ही {normalize_marathi_text(min_item['mar_name'])}ची संख्या आहे आणि ती सर्वात कमी आहे.<br>म्हणून {normalize_marathi_text(min_item['mar_name'])}, हे उत्तर."
    }
    questions_pool.append(q4)

    wrong_names_max = [x['eng_name'] for x in data if x != max_item]
    random.shuffle(wrong_names_max)
    q5 = {
        "q": f"Which {scenario['x_label_eng']} has maximum {scenario['y_label_eng']}?<br>#{scenario['y_label_mar']} सर्वात जास्त कोणत्या {scenario['x_label_mar']}चे आहे?",
        "corr": max_item['eng_name'],
        "wrong": wrong_names_max[:3],
        "sol": f"Ans : {max_item['eng_name']}.<br>From the given graph we can see that<br>{max_item['value']} is the value for {max_item['eng_name']} which is maximum.<br>Hence {max_item['eng_name']} is the answer.<br>#उत्तर : {normalize_marathi_text(max_item['mar_name'])}.<br>दिलेल्या स्तंभालेखानुसार <br>{max_item['value']} ही {normalize_marathi_text(max_item['mar_name'])}ची संख्या आहे आणि ती सर्वात जास्त आहे.<br>म्हणून {normalize_marathi_text(max_item['mar_name'])}, हे उत्तर."
    }
    questions_pool.append(q5)

    count = len(data)
    w_counts = [count - 1, count + 1, count + 2]
    item_names_eng = ", ".join([d['eng_name'] for d in data])
    item_names_mar = ", ".join([normalize_marathi_text(d['mar_name']) for d in data])
    q6 = {
        "q": f"How many different {scenario['x_label_eng']}s are used?<br>#एकूण किती वेगवेगळी {scenario['x_label_mar']} वापरली गेली आहेत?",
        "corr": str(count),
        "wrong": [str(w) for w in w_counts],
        "sol": f"Ans : {count}.<br>From the given graph we can see that, {item_names_eng} these are {count} {scenario['x_label_eng']}s used.<br>Hence {count} is the answer.<br>#उत्तर : {count}.<br>दिलेल्या स्तंभालेखानुसार {item_names_mar} अशी {count} {scenario['x_label_mar']} वापरली जातात हे उत्तर."
    }
    questions_pool.append(q6)

    q7 = {
        "q": f"Which are the different {scenario['x_label_eng']}s shown?<br>#दाखवलेली वेगवेगळी {scenario['x_label_mar']} कोणती आहेत?",
        "corr": item_names_eng,
        "wrong": ["A, B, C", "X, Y, Z", "Only " + data[0]['eng_name']],
        "sol": f"Ans : {item_names_eng}.<br>From the given graph, {scenario['x_label_eng']}s used are {item_names_eng} is the answer.<br>#उत्तर : {item_names_mar}.<br>दिलेल्या स्तंभालेखानुसार वापरलेली वेगवेगळी {scenario['x_label_mar']} {item_names_mar} अशी आहेत."
    }
    questions_pool.append(q7)

    w_val8 = [second_min['value'] + scale, second_min['value'] - scale, second_min['value'] + 5]
    q8 = {
        "q": f"How many {scenario['y_label_eng']} for second minimum used {scenario['x_label_eng']}?<br>#दोन क्रमांकाचे सर्वात कमी वापरले जाणाऱ्या {scenario['x_label_mar']}साठी {scenario['y_label_mar']} किती?",
        "corr": f"{second_min['eng_name']}, {second_min['value']}",
        "wrong": [f"{second_min['eng_name']}, {w}" for w in w_val8],
        "sol": f"Ans : {second_min['eng_name']}, {second_min['value']}.<br>From the given graph we can see that<br>Second minimum used {scenario['x_label_eng']} is {second_min['eng_name']} and value is {second_min['value']}.<br>#उत्तर : {normalize_marathi_text(second_min['mar_name'])}, {second_min['value']}.<br>दिलेल्या स्तंभालेखानुसार दोन क्रमांकाचे सर्वात कमी वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(second_min['mar_name'])} आहे, आणि त्याची संख्या {second_min['value']} हे उत्तर."
    }
    questions_pool.append(q8)

    w_val9 = [third_max['value'] + scale, third_max['value'] - scale, third_max['value'] + 5]
    q9 = {
        "q": f"How many {scenario['y_label_eng']} for third maximum used {scenario['x_label_eng']}?<br>#तीन क्रमांकाचे सर्वात जास्त वापरले जाणाऱ्या {scenario['x_label_mar']}साठी {scenario['y_label_mar']} किती?",
        "corr": f"{third_max['eng_name']}, {third_max['value']}",
        "wrong": [f"{third_max['eng_name']}, {w}" for w in w_val9],
        "sol": f"Ans : {third_max['eng_name']}, {third_max['value']}.<br>From the given graph we can see that<br>Third maximum used {scenario['x_label_eng']} is {third_max['eng_name']} and value is {third_max['value']}.<br>#उत्तर : {normalize_marathi_text(third_max['mar_name'])}, {third_max['value']}.<br>दिलेल्या स्तंभालेखानुसार तीन क्रमांकाचे सर्वात जास्त वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(third_max['mar_name'])} आहे, आणि त्याची संख्या {third_max['value']} हे उत्तर."
    }
    questions_pool.append(q9)

    w_diff = [diff_max_min + scale, diff_max_min - scale, diff_max_min + 10]
    q10 = {
        "q": f"How much is the difference between the {scenario['y_label_eng']}, for {scenario['x_label_eng']} used most and least?<br>#सर्वात जास्त आणि सर्वात कमी वापरले जाणारे {scenario['x_label_mar']}, यांच्या {scenario['y_label_mar']}मधील फरक किती आहे?",
        "corr": str(diff_max_min),
        "wrong": [str(w) for w in w_diff],
        "sol": f"Ans : {diff_max_min}.<br>From the given graph we can see that<br>Most used: {max_item['eng_name']} ({max_item['value']})<br>Least used: {min_item['eng_name']} ({min_item['value']})<br>By taking the difference of these two numbers we get <br>${max_item['value']} - {min_item['value']} = {diff_max_min}$ is the answer.<br>#उत्तर : {diff_max_min}.<br>दिलेल्या स्तंभालेखानुसार<br>सगळ्यात जास्त वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(max_item['mar_name'])} ({max_item['value']}) आहे.<br>सगळ्यात कमी वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(min_item['mar_name'])} ({min_item['value']}) आहे.<br>या दोन्ही संख्येतील फरक घेऊन आपल्याला <br>${max_item['value']} - {min_item['value']} = {diff_max_min}$ हे उत्तर मिळते."
    }
    questions_pool.append(q10)

    w_sum_mm = [sum_max_min + scale, sum_max_min - scale, sum_max_min + 10]
    q11 = {
        "q": f"How many are the total {scenario['y_label_eng']} for {scenario['x_label_eng']} used most and least?<br>#सर्वात जास्त आणि सर्वात कमी वापरले जाणारे {scenario['x_label_mar']}, यांच्या {scenario['y_label_mar']}ची एकूण संख्या किती आहे?",
        "corr": str(sum_max_min),
        "wrong": [str(w) for w in w_sum_mm],
        "sol": f"Ans : {sum_max_min}.<br>From the given graph we can see that<br>Most used: {max_item['eng_name']} ({max_item['value']})<br>Least used: {min_item['eng_name']} ({min_item['value']})<br>By taking the addition of these two numbers we get <br>${max_item['value']} + {min_item['value']} = {sum_max_min}$ is the answer.<br>#उत्तर : {sum_max_min}.<br>दिलेल्या स्तंभालेखानुसार<br>सगळ्यात जास्त वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(max_item['mar_name'])} ({max_item['value']}) आहे.<br>सगळ्यात कमी वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(min_item['mar_name'])} ({min_item['value']}) आहे.<br>या दोन्ही संख्यांची बेरीज करून आपल्याला <br>${max_item['value']} + {min_item['value']} = {sum_max_min}$ हे उत्तर मिळते."
    }
    questions_pool.append(q11)

    w_total = [total_val + scale, total_val - scale, total_val + 10]
    sum_str = " + ".join([str(x['value']) for x in data])
    q12 = {
        "q": f"How many are the total {scenario['y_label_eng']}?<br>#एकूण {scenario['y_label_mar']} किती आहेत?",
        "corr": str(total_val),
        "wrong": [str(w) for w in w_total],
        "sol": f"Ans : {total_val}.<br>By taking the total of all numbers we get <br>${sum_str} = {total_val}$ is the answer.<br>#उत्तर : {total_val}.<br>सर्व संख्यांची बेरीज करून <br>${sum_str} = {total_val}$ असे उत्तर मिळते."
    }
    questions_pool.append(q12)

    w_sum_23 = [sum_2nd_3rd_max + scale, sum_2nd_3rd_max - scale, sum_2nd_3rd_max + 10]
    q13 = {
        "q": f"How many are the total of {scenario['y_label_eng']} for second and third most used {scenario['x_label_eng']}?<br>#दोन क्रमांकाचे सर्वात जास्त आणि तीन क्रमांकाचे सर्वात जास्त वापरले जाणारे {scenario['x_label_mar']} यांची मिळून एकूण {scenario['y_label_mar']} किती?",
        "corr": str(sum_2nd_3rd_max),
        "wrong": [str(w) for w in w_sum_23],
        "sol": f"Ans : {sum_2nd_3rd_max}.<br>From the given graph we can see that<br>2nd most used: {second_max['eng_name']} ({second_max['value']})<br>3rd most used: {third_max['eng_name']} ({third_max['value']})<br>By taking the addition of these two numbers we get <br>${second_max['value']} + {third_max['value']} = {sum_2nd_3rd_max}$ is the answer.<br>#उत्तर : {sum_2nd_3rd_max}.<br>दिलेल्या स्तंभालेखानुसार <br>दोन क्रमांकाचे सर्वात जास्त वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(second_max['mar_name'])} ({second_max['value']}) आहे.<br>तीन क्रमांकाचे सर्वात जास्त वापरले जाणारे {scenario['x_label_mar']} {normalize_marathi_text(third_max['mar_name'])} ({third_max['value']}) आहे.<br>या दोन्ही संख्यांची बेरीज करून आपल्याला <br>${second_max['value']} + {third_max['value']} = {sum_2nd_3rd_max}$ हे उत्तर मिळते."
    }
    questions_pool.append(q13)

    selected_qs = random.sample(questions_pool, 6)

    for q_data in selected_qs:
        question_html = f"<p><b>Refer to the graph in the main question:</b><br>#<b>मुख्य प्रश्नातील आलेखाचा संदर्भ घ्या:</b><br>{q_data['q']}</p>"

        formatted_row = [
            current_sr_no, "", "", "09030201", question_html,
            q_data['corr'],
            "", "", "",
            q_data['wrong'][0], q_data['wrong'][1], q_data['wrong'][2],
            "", "", "", "",
            q_data['sol'], "", variation_num
        ]
        rows.append(formatted_row)
        current_sr_no += 1

    return rows, current_sr_no


def main():
    filename = "Intern_09030201_109_4_Assign9_Siddhi_Manjarekar.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")
    headers = [
        "Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
        "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
        "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3",
        "Time in seconds", "Difficulty Level", "Question (Image/ Audio/ Video)",
        "Contributor's Registered mailId", "Solution (Text Only)",
        "Solution (Image/ Audio/ Video)", "Variation Number"
    ]
    sheet.append(headers)
    try:
        num_groups = int(input("Enter number of question groups: "))
    except ValueError:
        num_groups = 50
    sr_no = 1
    for _ in range(num_groups):
        new_rows, next_sr_no = generate_question_group(sr_no, 109)
        for r in new_rows:
            sheet.append(r)
        # Added empty row after each group
        sheet.append([])
        sr_no = next_sr_no
    sheet.append(["****"])
    workbook.save(filename)
    print(f"Successfully generated {sr_no - 1} questions in {filename}")


if __name__ == "__main__":
    main()