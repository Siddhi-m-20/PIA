import random
import unicodedata
import openpyxl
import warnings
import os

warnings.filterwarnings('ignore')

def normalize_marathi_text(text):
    return unicodedata.normalize('NFC', text)

def get_weight_config_for_standard(standard):
    configs = {
        1: [(18, 3), (20, 3)],    
        2: [(20, 3), (22, 3)],    
        3: [(22, 3), (24, 3)],    
        4: [(24, 4), (26, 4)],    
        5: [(26, 4), (28, 4)],    
        6: [(28, 4), (30, 3)],    
        7: [(32, 4), (34, 4)],   
        8: [(35, 5), (38, 5)],   
        9: [(38, 5), (40, 5)],    
        10: [(40, 5), (42, 5)],   
        11: [(42, 6), (45, 6)],  
        12: [(45, 6), (48, 6)]    
    }
    return configs.get(standard, [(28, 4), (30, 3)]) 

def get_standard_name(standard):
    number_words_eng = {
        1: "First", 2: "Second", 3: "Third", 4: "Fourth", 5: "Fifth",
        6: "Sixth", 7: "Seventh", 8: "Eighth", 9: "Ninth", 10: "Tenth",
        11: "Eleventh", 12: "Twelfth"
    }
    
    number_words_mar = {
        1: "पहिलीच्या", 2: "दुसरीच्या", 3: "तिसरीच्या", 4: "चौथीच्या", 5: "पाचवीच्या",
        6: "सहावीच्या", 7: "सातवीच्या", 8: "आठवीच्या", 9: "नववीच्या", 10: "दहावीच्या",
        11: "अकरावीच्या", 12: "बारावीच्या"
    }
    
    return number_words_eng.get(standard, "Sixth"), number_words_mar.get(standard, "सहावीच्या")

def create_bar_graph_svg(heights, categories, scale=3, y_max=None):
    if y_max is None:
        y_max = max(15, max(heights) + 3)

    chart_height = 280
    bar_width = 60
    bar_spacing = 90
    start_x = 110
    margin_top = 30
    margin_bottom = 70
    margin_left = 90
    margin_right = 130 

    n = len(categories)
    content_width = start_x + (n - 1) * bar_spacing + bar_width
    svg_width = max(620, content_width + margin_right)
    svg_height = chart_height + margin_top + margin_bottom

    bar_heights = [(h / y_max) * chart_height for h in heights]
    bar_colors = ['royalblue', 'darkorange', 'darkgray', 'gold', 'forestgreen']

    svg = f'<svg width="{svg_width}" height="{svg_height}" xmlns="http://www.w3.org/2000/svg">'
    svg += '<rect width="100%" height="100%" fill="white"/>'

    for i in range(0, y_max + 1):
        y_pos = margin_top + (chart_height - (i / y_max) * chart_height)
        if i % 2 == 0:
            svg += f'<line x1="{margin_left}" y1="{y_pos}" x2="{svg_width - margin_right}" y2="{y_pos}" stroke="gray" stroke-width="1" opacity="0.8"/>'
            svg += f'<text x="{margin_left - 10}" y="{y_pos + 4}" text-anchor="end" font-size="13" font-family="Arial">{i}</text>'
        else:
            svg += f'<line x1="{margin_left}" y1="{y_pos}" x2="{svg_width - margin_right}" y2="{y_pos}" stroke="gray" stroke-width="0.5" opacity="0.6"/>'

    x_axis_y = margin_top + chart_height
    svg += f'<line x1="{margin_left}" y1="{margin_top}" x2="{margin_left}" y2="{x_axis_y}" stroke="black" stroke-width="1"/>'
    svg += f'<line x1="{margin_left}" y1="{x_axis_y}" x2="{svg_width - margin_right}" y2="{x_axis_y}" stroke="black" stroke-width="1"/>'

    for i, (height, color) in enumerate(zip(bar_heights, bar_colors)):
        x_pos = start_x + i * bar_spacing
        bar_y = margin_top + chart_height - height
        svg += f'<rect x="{x_pos}" y="{bar_y}" width="{bar_width}" height="{height}" fill="{color}" stroke="black" stroke-width="0.7"/>'

    label_y = margin_top + chart_height + 22
    for i, category in enumerate(categories):
        x_center = start_x + i * bar_spacing + (bar_width / 2)
        svg += f'<text x="{x_center}" y="{label_y}" text-anchor="middle" font-size="12" font-family="Arial" font-weight="bold">{category}</text>'

    eng_xlabel = "X-Axis"
    mar_xlabel = normalize_marathi_text("X-अक्ष")
    right_label_x = start_x + (n - 1) * bar_spacing + bar_width + 15
    right_label_y = margin_top + chart_height + 22
    svg += f'<text x="{right_label_x}" y="{right_label_y}" text-anchor="start" font-size="13" font-family="Arial">{eng_xlabel} / {mar_xlabel}</text>'

    weight_range_eng = "Weight Range (kg)"
    weight_range_mar = normalize_marathi_text("वजनी गट (किग्रॅ)")
    weight_text_y = margin_top + chart_height + 50
    weight_text_x = svg_width / 2
    svg += f'<text x="{weight_text_x}" y="{weight_text_y}" text-anchor="middle" font-size="16" font-family="Arial">{weight_range_eng} / {weight_range_mar}</text>'

    y_label_eng = "Number of Students"
    y_label_mar = normalize_marathi_text("विद्यार्थ्यांची संख्या")
    y_label_x = 45
    y_label_y = margin_top + (chart_height / 2)
    svg += f'<text x="{y_label_x}" y="{y_label_y}" text-anchor="middle" font-size="14" font-family="Arial" transform="rotate(-90 {y_label_x} {y_label_y})">{y_label_eng} / {y_label_mar}</text>'
    
    y_axis_text = "Y-Axis / Y-अक्ष"
    svg += f'<text x="{margin_left}" y="{margin_top - 10}" text-anchor="middle" font-size="13" font-family="Arial">{y_axis_text}</text>'

    svg += '</svg>'
    return svg

def create_bar_graph_html(heights, categories, scale, standard):
    chart_svg = create_bar_graph_svg(heights, categories, scale)
    if not chart_svg:
        return ""

    eng_std, mar_std = get_standard_name(standard)

    eng_title = f"Weight of Students in Class {standard}"
    mar_title = normalize_marathi_text(f"{mar_std} वर्गातील विद्यार्थ्यांचे वजन")
    eng_scale = f"Scale: 1 unit = {scale} students"
    mar_scale = normalize_marathi_text(f"प्रमाण: 1 एकक = {scale} विद्यार्थी")

    complete_chart = (
        f'<div style="text-align:center;border:1px solid &num;ccc;padding:12px;'
        f'max-width:900px;margin:0 auto;">'
        f'<div style="font-size:22px;font-weight:600">{eng_title} / {mar_title}</div>'
        f'<div style="font-size:14px;margin-top:4px;">({eng_scale} / {mar_scale})</div>'
        f'<div style="margin-top:10px;margin-bottom:8px;">'
        f'  {chart_svg}'
        f'</div>'
        f'</div>'
    )

    return complete_chart

def generate_weight_categories(standard, variety):
    configs = get_weight_config_for_standard(standard)
    start, step = configs[0] if variety == 1 else configs[1]
    categories = [f"{start + i*step}-{start + (i+1)*step}" for i in range(5)]
    return categories, start, step

def generate_student_counts(scale=3):
    heights = random.sample(range(1, 13), 5)
    
    student_counts = [h * scale for h in heights]
    
    return heights, student_counts

def create_solution_text(heights, student_counts, categories, max_category_idx, scale, start_weight, step):
    
    max_category = categories[max_category_idx]
    max_students = student_counts[max_category_idx]
    
    eng_sol = (
        f"Ans : Group ${max_category}$ kg and ${max_students}$ students.<br>"
        f"We are given the bar graph for the weight of all students.<br>"
        f"According to the weight range different groups are created. They are<br>"
    )
    
    for i in range(5):
        eng_sol += f"group ${i+1}$ - ${start_weight + i*step}$ kg upto ${start_weight + (i+1)*step}$ kg.<br>"
    
    eng_sol += (
        f"Scale used for this bar graph is $1$ unit $={scale}$ students.<br>"
        f"By knowing the scale factor and the height of each bar and by using the formula as,<br>"
        f"Number of students $=$ height of the bar $\\times$ scale.<br>"
        f"We get number of students in each group as,<br>"
    )
    
    for i in range(5):
        eng_sol += f"Number of students in group ${i+1}= {heights[i]} \\times {scale}= {student_counts[i]}$ students.<br>"
    
    eng_sol += (
        f"From this we can conclude that, highest number of students are in group ${max_category_idx+1}$ "
        f"with weight range as ${max_category}$ kg and their number is ${max_students}$.<br>"
        f"Hence, group ${max_category}$ kg and ${max_students}$ students is the answer.<br>&num;"
    )
    
    mar_sol = (
        f"उत्तर : गट ${max_category}$ किग्रॅ आणि ${max_students}$ विद्यार्थी.<br>"
        f"सर्व विद्यार्थ्यांचे वजन गटानुसार वेगवेगळे दाखविणारा स्तंभालेख दिला आहे.<br>"
        f"वजना नुसार बनविलेले वेगवेगळे गट खाली दाखविले आहेत.<br>"
    )
    
    for i in range(5):
        mar_sol += f"गट ${i+1} - {start_weight + i*step}$ किग्रॅ ते ${start_weight + (i+1)*step}$ किग्रॅ.<br>"
    
    mar_sol += (
        f"या स्तंभालेखासाठी वापरलेले प्रमाण $1$ एकक $={scale}$ विद्यार्थी असे आहे.<br>"
        f"या स्तंभालेखासाठी वापरलेले प्रमाण आणि यातील प्रत्येक स्तंभाची उंची माहिती असेल तर आपण खालील सूत्रानुसार प्रत्येक गटातील विद्यार्थ्यांची संख्या शोधू शकतो.<br>"
        f"विद्यार्थ्यांची संख्या  $=$ स्तंभाची उंची $\\times$ वापरलेले प्रमाण.<br>"
        f"या नुसार प्रत्येक गटातील विद्यार्थ्यांची संख्या खाली दिल्यानुसार आहे.<br>"
    )
    
    for i in range(5):
        mar_sol += f"गट ${i+1}$ मधील विद्यार्थ्यांची संख्या  $= {heights[i]} \\times {scale}= {student_counts[i]}$ विद्यार्थी.<br>"
    
    mar_sol += (
        f"या वरून आपण असा निष्कर्ष काढू शकतो की, सर्वात जास्त विद्यार्थी संख्या गट ${max_category_idx+1}$ म्हणजे "
        f"ज्या गटातील विद्यार्थ्यांचे वजन ${max_category}$ किग्रॅ या दरम्यान आहे आणि या गटातील विद्यार्थी संख्या ${max_students}$ आहे.<br>"
        f"म्हणून, गट ${max_category}$ किग्रॅ आणि ${max_students}$ विद्यार्थी हे उत्तर आहे.<br>"
    )
    
    return eng_sol + mar_sol

def generate_wrong_answers(correct_category, correct_count, categories, student_counts):
    
    wrong_answers = []
    used = set()
    used.add((correct_category, correct_count))
    
    attempts = 0
    max_attempts = 50
    
    while len(wrong_answers) < 3 and attempts < max_attempts:
        attempts += 1
        
        if len(wrong_answers) < 1:
            available_categories = [c for c in categories if c != correct_category]
            if available_categories:
                wrong_cat = random.choice(available_categories)
                if (wrong_cat, correct_count) not in used:
                    wrong_answers.append((wrong_cat, correct_count))
                    used.add((wrong_cat, correct_count))
                    continue
        
        if len(wrong_answers) < 2:
            other_counts = [c for c in student_counts if c != correct_count]
            if other_counts:
                wrong_count = random.choice(other_counts)
                if (correct_category, wrong_count) not in used:
                    wrong_answers.append((correct_category, wrong_count))
                    used.add((correct_category, wrong_count))
                    continue
        
        if len(wrong_answers) < 3:
            available_categories = [c for c in categories if c != correct_category]
            if available_categories:
                wrong_cat = random.choice(available_categories)
                wrong_count = correct_count + random.choice([-6, -3, 3, 6, 9])
                if wrong_count > 0 and (wrong_cat, wrong_count) not in used:
                    wrong_answers.append((wrong_cat, wrong_count))
                    used.add((wrong_cat, wrong_count))
                    continue
    
    return wrong_answers if len(wrong_answers) == 3 else None

def main():
    os.makedirs("09030201_106_2_Assign6_Siddhi_Manjarekar", exist_ok=True)
    filename = "09030201_106_2_Assign6_Siddhi_Manjarekar/Intern_09030201_106_2_Assign6_Siddhi_Manjarekar.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = ["Sr. No", "Question Type", "Answer Type", "Topic Number", "Question (Text Only)",
              "Correct Answer 1", "Correct Answer 2", "Correct Answer 3", "Correct Answer 4",
              "Wrong Answer 1", "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
              "Difficulty Level", "Question (Image/ Audio/ Video)", "Contributor's Registered mailId",
              "Solution (Text Only)", "Solution (Image/ Audio/ Video)", "Variation Number"]
    sheet.append(header)

    num_questions = int(input("How many questions do you want to generate? "))

    scale = 3
    generated_questions = {}
    generated_answers = {}
    generated_solutions = {}
    i = 1
    rows_to_add = []
    max_retries = num_questions * 10
    retry_count = 0

    while i <= num_questions and retry_count < max_retries:
        retry_count += 1
        
        standard = random.randint(1, 12)
        variety = 1 if i % 2 == 1 else 2
        
        categories, start_weight, step = generate_weight_categories(standard, variety)
        heights, student_counts = generate_student_counts(scale)
        
        max_count = max(student_counts)
        max_idx = student_counts.index(max_count)
        max_category = categories[max_idx]
        
        question_key = (tuple(sorted(heights)), variety, standard)
        if question_key in generated_questions:
            continue
        
        wrong_answers = generate_wrong_answers(max_category, max_count, categories, student_counts)
        if not wrong_answers:
            continue
        
        eng_std, mar_std = get_standard_name(standard)
        bar_graph_html = create_bar_graph_html(heights, categories, scale, standard)
        
        eng_question = (
            f"In class ${standard}$, weight of all students was measured. These measurements were displayed using bar graph as shown below. "
            f"In which of the weight category, maximum number of students appear and what is their number?<br>"
        )
        
        mar_question = normalize_marathi_text(
            f"{mar_std} वर्गातील सर्व विद्यार्थ्यांच्या वजनाची  नोंद केली.  ही माहिती स्तंभालेखाच्या आधारे दाखविली. "
            f"या स्तंभालेखाच्या साहाय्याने, कोणत्या वजन गटात सर्वात जास्त विद्यार्थी आहेत आणि त्यांची संख्या किती, हे सांगा.<br>"
        )
        
        question_text = f'<p>{eng_question}&num;{mar_question}{bar_graph_html}</p>'
        
        correct_answer = (
            f"Group ${max_category}$ kg and ${max_count}$ students.<br>&num;"
            f"गट ${max_category}$ किग्रॅ आणि ${max_count}$ विद्यार्थी.<br>"
        )
        
        wrong_answers_formatted = []
        for wrong_cat, wrong_count in wrong_answers:
            wrong_ans = (
                f"Group ${wrong_cat}$ kg and ${wrong_count}$ students.<br>&num;"
                f"गट ${wrong_cat}$ किग्रॅ आणि ${wrong_count}$ विद्यार्थी.<br>"
            )
            wrong_answers_formatted.append(wrong_ans)
        
        if correct_answer in wrong_answers_formatted:
            continue
        if len(set(wrong_answers_formatted)) != 3:
            continue
        
        solution = create_solution_text(heights, student_counts, categories, max_idx, scale, start_weight, step)
        
        if question_text in generated_questions:
            continue
        if correct_answer in generated_answers:
            continue
        if solution in generated_solutions:
            continue
        
        generated_questions[question_text] = i
        generated_answers[correct_answer] = i
        generated_solutions[solution] = i
        
        row_data = [i, "Text", 1, "09030201", question_text, correct_answer,
                    "", "", "", wrong_answers_formatted[0], wrong_answers_formatted[1], 
                    wrong_answers_formatted[2], 120, 2, "", "siddhimanjarekar2@gmail.com", 
                    solution, "", 106]
        
        rows_to_add.append(row_data)
        i += 1
    
    for row in rows_to_add:
        sheet.append(row)
    
    sheet.append(["****"])
    
    workbook.save(filename)
    print(f"Successfully generated {len(rows_to_add)} questions in '{filename}'!")

if __name__ == "__main__":
    main()