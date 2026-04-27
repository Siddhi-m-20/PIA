import random
import unicodedata
import openpyxl

def normalize_marathi_text(text):
    return unicodedata.normalize("NFC", text)

def get_mistake_descriptions():
    return {
        1: {
            "eng": "Scale used in the graph and scale mentioned is not identical",
            "mar": normalize_marathi_text("स्तंभालेखात वापरलेले प्रमाण आणि नमूद केलेले प्रमाण एकच नाही")
        },
        2: {
            "eng": "Width of all bars is not uniform",
            "mar": normalize_marathi_text("सर्व स्तंभाची रुंदी एकसमान नाही")
        },
        3: {
            "eng": "Scale shown on Y-axis is not uniform",
            "mar": normalize_marathi_text("Y-अक्षावरील प्रमाण एकसमान नाही")
        },
        4: {
            "eng": "Value written on top of the bars must match their height as per the scale",
            "mar": normalize_marathi_text("प्रत्येक स्तंभाच्या डोक्यावर लिहिलेली किंमत दिलेल्या प्रमाणानुसार स्तंभाच्या उंचीशी जुळली पाहिजे")
        },
        5: {
            "eng": "Bar for one day is missing",
            "mar": normalize_marathi_text("एका दिवसाचा स्तंभ गहाळ आहे")
        }
    }

def create_bar_graph_svg_with_mistake(days, values, mistake_type):
    chart_height = 280
    bar_width = 60
    bar_spacing = 90
    start_x = 110
    margin_top = 50
    margin_bottom = 70
    margin_left = 90
    margin_right = 130
    n = len(days)
    mistake_bar_idx = random.randint(0, n - 1)
    y_max = max(values)
    rounded_ymax = max(((y_max // 100) + 2) * 100, 500)
    content_width = start_x + (n - 1) * bar_spacing + bar_width
    svg_width = max(620, content_width + margin_right)
    svg_height = chart_height + margin_top + margin_bottom
    svg = f"<svg width='{svg_width}' height='{svg_height}' xmlns='http://www.w3.org/2000/svg'>"
    svg += "<rect width='100%' height='100%' fill='white'/>"
    if mistake_type == 3:
        custom_labels = [0,150, 200, 450, 700,1000]

        num_ticks = len(custom_labels)
        tick_positions = [
            margin_top + chart_height - (i / (num_ticks - 1)) * chart_height
            for i in range(num_ticks)
        ]
        y_max_for_plot = custom_labels[-1]
        adjusted_values = []
        for v in values:
            if v >= y_max_for_plot:
                v = y_max_for_plot - 200
            adjusted_values.append(v)

        bar_heights = [(v / y_max_for_plot) * chart_height for v in adjusted_values]

        for i, (label, pos) in enumerate(zip(custom_labels, tick_positions)):
            svg += (
                f"<line x1='{margin_left}' y1='{pos}' "
                f"x2='{svg_width - margin_right}' y2='{pos}' "
                f"stroke='gray' opacity='0.8'/>"
            )
            svg += (
                f"<text x='{margin_left - 10}' y='{pos + 4}' "
                f"text-anchor='end'>{label}</text>"
            )
            if i < num_ticks - 1:
                next_pos = tick_positions[i + 1]
                minor_gap = (pos - next_pos) / 5
                for j in range(1, 5):
                    mpos = pos - j * minor_gap
                    svg += (
                        f"<line x1='{margin_left}' y1='{mpos}' "
                        f"x2='{svg_width - margin_right}' y2='{mpos}' "
                        f"stroke='lightgray' opacity='0.5'/>"
                    )
    else:
        num_gridlines = 6
        bar_heights = [(v / rounded_ymax) * chart_height for v in values]
        for i in range(num_gridlines):
            y_val = int((i / (num_gridlines - 1)) * rounded_ymax)
            y_pos = margin_top + chart_height - (i / (num_gridlines - 1)) * chart_height
            svg += f"<line x1='{margin_left}' y1='{y_pos}' x2='{svg_width - margin_right}' y2='{y_pos}' stroke='gray' opacity='0.8'/>"
            svg += f"<text x='{margin_left - 10}' y='{y_pos + 4}' text-anchor='end'>{y_val}</text>"
            if i < num_gridlines - 1:
                next_y = margin_top + chart_height - ((i + 1) / (num_gridlines - 1)) * chart_height
                minor_gap = (y_pos - next_y) / 5
                for j in range(1, 5):
                    mpos = y_pos - j * minor_gap
                    svg += f"<line x1='{margin_left}' y1='{mpos}' x2='{svg_width - margin_right}' y2='{mpos}' stroke='lightgray' opacity='0.5'/>"
    x_axis_y = margin_top + chart_height
    svg += f"<line x1='{margin_left}' y1='{margin_top}' x2='{margin_left}' y2='{x_axis_y}' stroke='black'/>"
    svg += f"<line x1='{margin_left}' y1='{x_axis_y}' x2='{svg_width - margin_right}' y2='{x_axis_y}' stroke='black'/>"
    svg += f"<text x='{margin_left}' y='{margin_top - 15}' text-anchor='middle' font-size='13' font-family='Arial'>Y-Axis / Y-अक्ष</text>"
    bar_colors = ["blue", "orange", "gray", "gold", "green"]
    wrong_value_for_mistake4 = None
    bar_idx = 0
    for i, ((day_eng, day_mar), val, h, color) in enumerate(
        zip(days, values, bar_heights, bar_colors)
    ):
        if mistake_type == 5 and i == mistake_bar_idx:
            continue
        x_pos = start_x + bar_idx * bar_spacing
        current_width = bar_width
        current_height = h
        if mistake_type == 2 and i == mistake_bar_idx:
            current_width = int(bar_width * 1.6)
        bar_y = margin_top + chart_height - current_height
        svg += f"<rect x='{x_pos}' y='{bar_y}' width='{current_width}' height='{current_height}' fill='{color}' stroke='black'/>"
        display_val = val
        if mistake_type == 4 and i == mistake_bar_idx:
            display_val = val + random.choice([100, -150, 200, -100])
            wrong_value_for_mistake4 = display_val
        svg += f"<text x='{x_pos + current_width/2}' y='{bar_y - 5}' text-anchor='middle'>{display_val}</text>"
        bar_idx += 1
    label_y_eng = margin_top + chart_height + 18
    label_y_mar = label_y_eng + 15
    bar_idx = 0
    for i, (eng, mar) in enumerate(days):
        if mistake_type == 5 and i == mistake_bar_idx:
            continue
        xc = start_x + bar_idx * bar_spacing + bar_width / 2
        svg += f"<text x='{xc}' y='{label_y_eng}' text-anchor='middle' font-size='14' font-family='Arial'>{eng}</text>"
        svg += f"<text x='{xc}' y='{label_y_mar}' text-anchor='middle' font-size='14' font-family='Arial'>{mar}</text>"
        bar_idx += 1
    if mistake_type == 1:
        wrong_scale = random.choice([50, 200, 300])
        svg += (
            f"<text x='{svg_width/2}' y='{margin_top-15}' text-anchor='middle'>"
            f"Scale: 1 unit = ₹{wrong_scale} / प्रमाण: 1 एकक = ₹{wrong_scale}</text>"
        )
    elif mistake_type == 3:
        svg += (
            f"<text x='{svg_width/2}' y='{margin_top-15}' text-anchor='middle'>"
            f"Scale: 1 unit = ₹150 / प्रमाण: 1 एकक = ₹150</text>"
        )
    else:
        scale = rounded_ymax // 5
        svg += (
            f"<text x='{svg_width/2}' y='{margin_top-15}' text-anchor='middle'>"
            f"Scale: 1 unit = ₹{scale} / प्रमाण: 1 एकक = ₹{scale}</text>"
        )
    svg += f"<text x='{svg_width - margin_right + 15}' y='{label_y_mar}' text-anchor='start' font-size='13' font-family='Arial'>X-Axis / X-अक्ष</text>"
    svg += f"<text x='{svg_width / 2}' y='{svg_height - 20}' text-anchor='middle' font-size='14' font-family='Arial'>Days / दिवस</text>"
    svg += f"<text x='35' y='{margin_top + chart_height / 2}' transform='rotate(-90 35 {margin_top + chart_height / 2})' text-anchor='middle' font-size='14' font-family='Arial'>Sale (₹) / विक्री (₹)</text>"
    svg += "</svg>"
    svg = svg.replace("#", "&num;")
    return svg, mistake_bar_idx, wrong_value_for_mistake4

def generate_wrong_answers(correct_mistake_type):
    mistakes = get_mistake_descriptions()
    wrong = []
    for mtype, info in mistakes.items():
        if mtype == correct_mistake_type:
            continue
        wrong.append(f"{info['eng']}.<br>&num;{info['mar']}.<br>")
    random.shuffle(wrong)
    return wrong[:3]

def create_solution_text(mistake_type, days, values, mistake_bar_idx, wrong_value_for_mistake4=None):
    mistakes = get_mistake_descriptions()
    correct_desc = mistakes[mistake_type]
    cond_eng = [
        "Data for each day under consideration is present",
        "Scale shown is uniform along Y-axis and height of each bar is in accordance with the scale",
        "Width of each bar is identical and spacing between bars is also identical",
        "Value written on top of the bars must match their height as per the scale",
        "Scale used in the graph and scale mentioned is identical"
    ]
    cond_mar = [
        normalize_marathi_text("दिलेल्या प्रत्येक दिवसाची माहिती स्तंभालेखात आहे"),
        normalize_marathi_text("Y-अक्षावरील प्रमाण एकसमान आहे आणि प्रत्येक स्तंभाची उंची प्रमाणा नुसार आहे"),
        normalize_marathi_text("सर्व स्तंभांची रुंदी समान आहे व लगतच्या दोन स्तंभांमधील अंतर समान आहे"),
        normalize_marathi_text("प्रत्येक स्तंभाच्या डोक्यावर लिहिलेली किंमत दिलेल्या प्रमाणानुसार स्तंभाच्या उंचीशी जुळली पाहिजे"),
        normalize_marathi_text("स्तंभालेखात वापरलेले प्रमाण आणि नमूद केलेले प्रमाण एकच आहे")
    ]
    if mistake_type == 1:
        failed_idx = 4
    elif mistake_type == 2:
        failed_idx = 2
    elif mistake_type == 3:
        failed_idx = 1
    elif mistake_type == 4:
        failed_idx = 3
    elif mistake_type == 5:
        failed_idx = 0

    failed_cond_eng = cond_eng[failed_idx]
    failed_cond_mar = cond_mar[failed_idx]

    eng_solution = (
        f"Ans : {correct_desc['eng']}.<br>"
        "To identify the mistake in the bar graph, we will check the graph for following conditions,<br>"
        f"$1.$ {cond_eng[0]}.<br>"
        f"$2.$ {cond_eng[1]}.<br>"
        f"$3.$ {cond_eng[2]}.<br>"
        f"$4.$ {cond_eng[3]}.<br>"
        f"$5.$ {cond_eng[4]}.<br>"
        f"By checking the given graph for compliance of these conditions, we can observe that, <b>{failed_cond_eng}</b> is <b>not complied</b>.<br>"
        "The given graph fulfills all other conditions.<br>"
        f"Hence, <b>{correct_desc['eng']}</b> is the answer.<br>"
    )
    mar_solution = (
        f"उत्तर : {correct_desc['mar']}.<br>"
        "दिलेल्या स्तंभालेखात काय चूक आहे, हे शोधण्यासाठी, आपल्याला पुढील बाबी पडताळून पाहू.<br>"
        f"$1.$ {cond_mar[0]}.<br>"
        f"$2.$ {cond_mar[1]}.<br>"
        f"$3.$ {cond_mar[2]}.<br>"
        f"$4.$ {cond_mar[3]}.<br>"
        f"$5.$ {cond_mar[4]}.<br>"
        f"या नुसार दिलेल्या सर्व अटी पडताळून पाहता <b>{failed_cond_mar}</b> ही एकच <b>अट पाळली गेली नाही</b> असे लक्षात येते. <br>"
        "या शिवाय उरलेल्या सर्व अटींची पूर्तता होते आहे. <br>"
        f"म्हणून, <b>{correct_desc['mar']}</b> हे उत्तर.<br>"
    )
    return f"{eng_solution}&num;{mar_solution}"

def main():
    filename = "Intern_09030201_107_2_Assign7_Siddhi_Manjarekar.xlsx"
    workbook = openpyxl.Workbook()

    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])

    workbook.create_sheet("Instruction")
    sheet = workbook.create_sheet("Questions")

    header = [
        "Sr. No", "Question Type", "Answer Type", "Topic Number",
        "Question (Text Only)", "Correct Answer 1", "Correct Answer 2",
        "Correct Answer 3", "Correct Answer 4", "Wrong Answer 1",
        "Wrong Answer 2", "Wrong Answer 3", "Time in seconds",
        "Difficulty Level", "Question (Image/ Audio/ Video)",
        "Contributor's Registered mailId", "Solution (Text Only)",
        "Solution (Image/ Audio/ Video)", "Variation Number"
    ]

    sheet.append(header)

    try:
        num_questions = int(input("How many questions do you want to generate? ").strip())
        if num_questions <= 0:
            num_questions = 5
    except Exception:
        num_questions = 5

    days = [
        ("Monday", "सोमवार"),
        ("Tuesday", "मंगळवार"),
        ("Wednesday", "बुधवार"),
        ("Thursday", "गुरुवार"),
        ("Friday", "शुक्रवार")
    ]

    generated = set()
    rows = []
    i = 1

    allowed_values = [100, 150, 200, 250, 300, 350, 400, 450,
                      500, 550, 600, 650, 700, 750, 800]

    while i <= num_questions:
        values = [random.choice(allowed_values) for _ in range(5)]
        mistake_type = random.randint(1, 5)

        svg, idx, wrong_val = create_bar_graph_svg_with_mistake(days, values, mistake_type)

        graph_html = (
            f"<div style='text-align:center;border:1px solid &num;ccc;padding:12px;"
            f"max-width:900px;margin:0 auto;'>"
            f"<div style='font-size:18px;font-weight:600;margin-bottom:10px;'>"
            f"Daily Sale in Shop (Monday to Friday)<br>"
            f"{normalize_marathi_text('दुकानातील दैनिक विक्री (सोमवार ते शुक्रवार)')}</div>"
            f"<div style='margin-top:10px;margin-bottom:8px;'>{svg}</div>"
            f"</div>"
        )

        eng_q = (
            "A bar graph is as shown below with correct values mentioned on the top of each bar.<br>"
            "It shows daily sale in a shop from Monday to Friday.<br>"
            "Something is wrong in this bar graph.<br>"
            "The mistake in the graph is ______________.<br>"
        )

        mar_q = normalize_marathi_text(
            "खाली एक स्तंभालेख दिला आहे आणि यात प्रत्येक स्तंभाच्या डोक्यावर लिहिलेली किंमत बरोबर आहे.<br>"
            "हा स्तंभालेख, सोमवार ते शुक्रवार या दिवसात झालेली दुकानातील दैनिक विक्री दाखवतो.<br>"
            "या स्तंभालेखात काहीतरी चुकीचे आहे.<br>"
            "स्तंभालेखातील चूक ______________ अशी आहे.<br>"
        )

        question = f"<p>{eng_q}&num;{mar_q}{graph_html}</p>"

        key = (tuple(values), mistake_type)
        if key in generated:
            continue

        mistakes = get_mistake_descriptions()
        correct_ans = (
            f"{mistakes[mistake_type]['eng']}.<br>&num;{mistakes[mistake_type]['mar']}.<br>"
        )

        wrong_options = generate_wrong_answers(mistake_type)

        solution = create_solution_text(mistake_type, days, values, idx, wrong_val)

        row = [
            i, "Text", 1, "09030201", question,
            correct_ans, "", "", "", wrong_options[0],
            wrong_options[1], wrong_options[2], 180, 2, "",
            "siddhimanjarekar274@gmail.com", solution, "", 107
        ]

        rows.append(row)
        generated.add(key)
        i += 1

    for r in rows:
        sheet.append(r)

    sheet.append(["****"])
    workbook.save(filename)
    print(f"Saved {len(rows)} questions → {filename}")

if __name__ == "__main__":
    main()